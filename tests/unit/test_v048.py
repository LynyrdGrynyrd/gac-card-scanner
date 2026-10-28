"""v0.4.8 regression tests (bot reports qa/*-v047.md). Made-up cards only."""
import io
import json
from test_app import mk


def _appr(ui, st, batch, **f):
    a = mk(ui, st, **f); ui.set_meta(a, batch=batch, batch_label=batch.upper()); st.update(a, status="approved", reviewer="RR"); return a


def _exp_rows(r):
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    return [[c.value for c in row] for row in ws.iter_rows(min_row=2) if any(c.value for c in row)]


# ---- #1 HIGH: an export that leaves out an edited-after-export card makes no empty file and marks nothing
def test_v048_edited_after_export_alone_gives_no_file(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v48a", first_name="Edda", last_name="Twice", company="Twice Co", emails=["edda@v48a.example.com"])
    r = adm.get("/api/export/xlsx?status=approved&batch=v48a"); assert r.status_code == 200 and "spreadsheet" in r.headers["content-type"]
    assert len(_exp_rows(r)) == 1 and r.headers["X-Card-Count"] == "1"
    # an edit after the export
    assert adm.post(f"/api/cards/{a}", json={"fields": {"title": "Boss"}}).status_code == 200
    m = ui.meta(a); assert m.get("reexport")
    n_before = ui.db().execute("SELECT COUNT(*) FROM exports").fetchone()[0]
    h_before = len([h for h in ui.history(a) if h.get("action") == "exported"])
    ck = adm.get("/api/export/check?status=approved&batch=v48a").json()
    assert ck["rows"] == 0 and ck["held"] == 1                      # the page can say "Nothing new to import."
    r2 = adm.get("/api/export/xlsx?status=approved&batch=v48a"); assert r2.status_code == 200
    j = r2.json(); assert j["nothing"] and j["message"] == "Nothing new to import." and a in j["changed_ids"]
    assert "1 edited after export" in j["ready"]["text"]
    assert ui.db().execute("SELECT COUNT(*) FROM exports").fetchone()[0] == n_before   # no export row, no file
    assert ui.meta(a).get("reexport")                            # the warning stays
    assert len([h for h in ui.history(a) if h.get("action") == "exported"]) == h_before
    # with a second, clean card the file holds only the real row, and only that card is marked
    b = _appr(ui, st, "v48a", first_name="Finn", last_name="Fresh", company="Fresh Co", emails=["finn@v48a.example.com"])
    r3 = adm.get("/api/export/xlsx?status=approved&batch=v48a"); rows = _exp_rows(r3)
    assert len(rows) == 1 and "Finn" in json.dumps(rows, default=str) and r3.headers["X-Card-Count"] == "1"
    eid = r3.headers["X-Export-Id"]; row = ui.db().execute("SELECT card_ids,count FROM exports WHERE id=?", (eid,)).fetchone()
    assert json.loads(row[0]) == [b] and row[1] == 1
    assert ui.meta(a).get("reexport")
    rd = adm.get(f"/api/exports/{eid}/companion/readme.txt").text
    assert f"#{a}: Edited after export, so it is not in this file" in rd and "by hand" not in rd
    # the admin tick includes it anyway, and only then is it marked again
    r4 = adm.get("/api/export/xlsx?status=approved&batch=v48a&include_changed=1"); assert len(_exp_rows(r4)) == 2   # the one tick includes repeats too
    assert not ui.meta(a).get("reexport")


# ---- #2 HIGH: "All approved" does not re-add cards already in an earlier IMPORT-ME file
def test_v048_all_approved_skips_earlier_import_rows(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v48b", first_name="Gus", last_name="Repeat", company="Repeat Ltd", emails=["gus@v48b.example.com"])
    assert len(_exp_rows(adm.get("/api/export/xlsx?status=approved_new&batch=v48b"))) == 1
    # the same person scanned again in another batch: already in MC Trade from the earlier file
    a2 = _appr(ui, st, "v48c", first_name="Gus", last_name="Repeat", company="Repeat Ltd", emails=[])
    c = _appr(ui, st, "v48c", first_name="Hal", last_name="New", company="New Ltd", emails=["hal@v48c.example.com"])
    p = adm.get("/api/export/xlsx?status=approved&preview=1")
    s = json.loads(p.headers["X-Summary"])
    assert a in s["repeat_ids"] and a2 in s["repeat_ids"] and c not in s["repeat_ids"]
    assert s["decisions"][str(a)] != "new" and s["decisions"][str(c)] == "new"
    assert "already exported" in s["ready"]["text"]
    r = adm.get("/api/export/xlsx?status=approved&batch=v48c"); rows = json.dumps(_exp_rows(r), default=str)
    assert "Hal" in rows and "Gus" not in rows
    rd = adm.get(f"/api/exports/{r.headers['X-Export-Id']}/companion/readme.txt").text
    assert f"#{a2}: Already exported, so not in this file" in rd
    # admin tick: include anyway, with a warning
    r2 = adm.get("/api/export/xlsx?status=approved&batch=v48c&include_changed=1&preview=1")
    s2 = json.loads(r2.headers["X-Summary"]); assert s2["decisions"][str(a2)] == "new"


def test_v048_csv_only_card_is_new_with_a_plain_note(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v48d", first_name="Ivy", last_name="Csv", company="Csv Inc", emails=["ivy@v48d.example.com"])
    assert adm.get("/api/export/csv?status=approved&batch=v48d").status_code == 200   # contacts list only
    p = adm.get("/api/export/xlsx?status=approved&batch=v48d&preview=1"); s = json.loads(p.headers["X-Summary"])
    assert s["decisions"][str(a)] == "new" and a in s["csv_only_ids"]
    ex = adm.get("/api/export/xlsx?status=approved&batch=v48d&preview=1")
    from app import exporting as E
    assert E.CSV_ONLY == "Sent before only in a contacts list (CSV), not to MC Trade."


def test_v048_ready_line_and_totals_plain():
    from app import exporting as E
    assert E.ready_line({"rows": 0, "clean": 0, "to_check": 0, "changed": 1}) == "Nothing new to import. Not in the file: 1 edited after export."
    assert E.ready_line({"rows": 2, "clean": 2, "to_check": 0, "repeat": 1, "waiting": 1}).startswith("Ready to import: 2 rows (2 clean, 0 to check). Not in the file: ")
    from app import mctrade_export as MX
    cnt = MX.counts({1: (MX.NEW, ""), 2: (MX.NEW, ""), 3: (MX.NEW, "")})
    t = E.readme_totals(cnt, to_check=1)
    assert t == "In the file: 3 new contacts (1 to check)." and "·" not in t and "of the rows" not in t


# ---- #3: the save-time warning on an already-exported card
def test_v048_save_notice_on_exported_card(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v48e", first_name="Jo", last_name="Saved", company="Saved Co", emails=["jo@v48e.example.com"])
    adm.get("/api/export/xlsx?status=approved&batch=v48e")
    r = adm.post(f"/api/cards/{a}", json={"fields": {"title": "Chief"}}).json()
    assert r["notice"] == "Saved. This card was exported before, so this edit won't be exported again unless an admin includes it."
    src = "\n".join(open(f"app/{n}", encoding="utf-8").read() for n in ("ui.py", "exporting.py")) + open("app/static/ui/views/export.js", encoding="utf-8").read()
    assert "MC Trade by hand" not in src and "Changed after export: export it again" not in src


# ---- #4 MED: the reviewer's CSV Check column and History detail data name only same-batch cards
def test_v048_reviewer_csv_and_history_detail_same_batch_only(env):
    from test_v047 import _world
    ui, st = env["ui"], env["store"]
    r, o, ids = _world(env, "v48p")
    cur, oth, old = ids["cur"], ids["oth"], ids["old"]
    st.update(ids["cur2"], status="rejected", reviewer="RR")
    st.update(cur, status="approved", reviewer="RR"); st.update(oth, status="approved", reviewer="OO"); st.update(old, status="approved", reviewer="RR")
    ui.log(cur, "RR", "cross_batch_confirmed", detail={"note": "approved although an approved look-alike exists in another batch", "cards": [oth, old, ids["cur2"]]})
    ui.log(cur, "RR", "reexport_needed", detail={"why": f"details merged from #{oth}"})
    ui.log(cur, "RR", "undone", detail={"what": f"Merge card #{old} into #{cur}", "target": oth})
    rc = r.get("/api/export/csv?status=approved&batch=v48pcur"); assert rc.status_code == 200, rc.text[:400]
    line = [x for x in rc.text.splitlines() if x.startswith(f"{cur},")][0]
    assert f"#{oth}" not in line and f"#{old}" not in line and f"#{ids['cur2']}" in line and "in other batches" in line
    pv = r.get("/api/export/xlsx?status=approved&batch=v48pcur&preview=1&view=1").json()
    assert f"#{oth}" not in json.dumps(pv.get("warnings")) and f"#{old}" not in json.dumps(pv.get("warnings"))
    h = json.dumps(r.get(f"/api/cards/{cur}").json()["history"])
    assert f"#{oth}" not in h and f"#{old}" not in h and f'"target": {oth}' not in h
    assert f"#{oth}" in json.dumps(env["adm"].get(f"/api/cards/{cur}").json()["history"])   # an admin still sees them


def test_v048_9_misreads_soft_and_hard():
    from app import ui, plausibility as P
    base = {"first_name": "Ann", "last_name": "Lee", "emails": ["a@b.example.com"], "phones": []}
    assert not ui.company_soft_hint(dict(base, company="E2open"))
    assert "company" not in ui.hard_errors(dict(base, company="4ward Logistics"))
    assert "title" not in ui.hard_errors(dict(base, title="Owner|Founder"))
    assert "title" in ui.hard_errors(dict(base, title="Vlce President"))
    assert "title" in ui.hard_errors(dict(base, title="Managr"))
    assert "name" in ui.hard_errors(dict(base, first_name="Rnbert"))
    assert "title" in ui.hard_errors(dict(base, title="Manag|ng Director"))
    assert not P.digit_brand_ok("4nalytics")


# ---- #11 LOW: the admin's All-view undo keeps to the batch of their newest action; the button names the action
def test_v048_admin_undo_all_view_batch_rule(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = mk(ui, st, first_name="Olga", last_name="Older", emails=["oo@v48u.example.com"]); ui.set_meta(a, batch="v48u1", batch_label="Older batch · Oct 28 · AD")
    b = mk(ui, st, first_name="Nina", last_name="Newer", emails=["nn@v48u.example.com"]); ui.set_meta(b, batch="v48u2", batch_label="Newer batch · Oct 28 · AD")
    assert adm.post(f"/api/cards/{a}", json={"fields": {"title": "Undo A"}}).status_code == 200
    assert adm.post(f"/api/cards/{b}", json={"fields": {"title": "Undo B"}}).status_code == 200
    p = adm.get("/api/undo").json()
    assert p["batch"] == "v48u2" and p["batch_short"] == "Newer batch" and p["short"].startswith("Edit #") and p["window_min"] == 120
    old = ui.db().execute("SELECT id FROM actions WHERE label LIKE ? ORDER BY id DESC LIMIT 1", (f"%#{a}%",)).fetchone()[0]
    r = adm.post("/api/undo", json={"id": old}); assert r.status_code == 404   # the older batch's action is not reachable from All
    assert ui.cards_get(a)["fields"]["title"] == "Undo A" if hasattr(ui, "cards_get") else True
    assert adm.post("/api/undo", json={"id": old, "batch": "v48u1"}).status_code == 200


# ---- #12 LOW: an upload with your own existing batch name (no id) goes into that batch, no "(2)"
def test_v048_upload_label_reuses_own_batch(env, monkeypatch):
    ui, st = env["ui"], env["store"]; seen = []
    async def fake(data, fname, **k): seen.append(k.get("batch")); return {"job": "x"}
    monkeypatch.setattr(ui, "enqueue_upload", fake)
    from test_v047 import _user, _uid, _ini
    rc = _user(env, "up48"); ui.claim_batch("v48-own-batch", {"id": _uid("up48")})
    c = mk(ui, st, first_name="Ulla", last_name="Up", emails=["uu@v48up.example.com"], by=_ini("up48")); ui.set_meta(c, batch="v48-own-batch", batch_label="Spring Fair 48")
    assert rc.post("/api/jobs", data={"batch_label": "spring fair 48"}, files={"image": ("a.jpg", b"\xff\xd8\xff\xd9", "image/jpeg")}).status_code == 200
    assert seen[-1] == "v48-own-batch"
    # someone else's batch with that name is never reused
    rc2 = _user(env, "up48b")
    assert rc2.post("/api/jobs", data={"batch_label": "Spring Fair 48"}, files={"image": ("a.jpg", b"\xff\xd8\xff\xd9", "image/jpeg")}).status_code == 200
    assert seen[-1] != "v48-own-batch"


# ---- #5 MED: "Exported before" only for the same person, and it names which export
def test_v048_exported_before_only_when_relevant(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v48x", first_name="Petra", last_name="Same", company="Same Co", emails=["petra@v48x.example.com"])
    assert adm.get("/api/export/xlsx?status=approved&batch=v48x").status_code == 200
    b = mk(ui, st, first_name="Petra", last_name="Same", company="Same Co", emails=["petra@v48x.example.com"]); ui.set_meta(b, batch="v48y")
    c = mk(ui, st, first_name="Petra", last_name="Samel", company="Other Firm", emails=["ps@other48.example.com"]); ui.set_meta(c, batch="v48y")
    xb = adm.get(f"/api/cards/{b}").json().get("exported_elsewhere") or []
    assert any(x["card_id"] == a and x.get("last_ts") for x in xb)
    xc = adm.get(f"/api/cards/{c}").json().get("exported_elsewhere") or []
    assert not any(x.get("card_id") == a for x in xc)   # a similar name alone is not "exported before"


# ---- #6 MED: Needs admin: one audit row; the flag clears when nothing is left to confirm (also old flags, at start)
def test_v048_needs_admin_single_row_and_settle(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = mk(ui, st, first_name="Nell", last_name="Admin", title="Managr", emails=["na@v48n.example.com"]); ui.set_meta(a, batch="v48n")
    assert adm.post(f"/api/cards/{a}/needs-admin", json={}).json()["needs_admin"]
    n_rows = [h for h in ui.history(a) if h.get("action") == "needs_admin"]; assert len(n_rows) == 1
    assert adm.post(f"/api/cards/{a}", json={"fields": {"title": "Manager"}}).status_code == 200
    assert not ui.meta(a).get("needs_admin")
    # an old (v0.4.6) flag with nothing left to confirm comes off in the start-up sweep
    b = mk(ui, st, first_name="Olly", last_name="Oldflag", title="Owner", emails=["oo@v48n.example.com"]); ui.set_meta(b, batch="v48n")
    ui.set_meta(b, needs_admin=json.dumps({"reason": "old", "by": "XX", "ts": 1}))
    assert ui.na_sweep() >= 1 and not ui.meta(b).get("needs_admin")


# ---- #7 MED: a reviewer's badge = pending cards of the one batch the queue shows
def test_v048_badge_one_batch(env):
    ui, st = env["ui"], env["store"]
    from test_v047 import _user, _uid, _ini
    rc = _user(env, "bdg48"); u = {"id": _uid("bdg48")}; ui.claim_batch("v48g1", u); ui.claim_batch("v48g2", u)
    for b_, n_ in (("v48g1", 3), ("v48g2", 1)):
        for k in range(n_): i = mk(ui, st, first_name=f"G{b_}{k}", last_name="Badge", emails=[f"g{k}@{b_}.example.com"], by=_ini("bdg48")); ui.set_meta(i, batch=b_, batch_label=b_)
    c = rc.get("/api/counts").json()
    assert c["badge_batch"] == "v48g2" and c["mine_pending"] == 1 and c["own_pending"] == 4
    assert rc.get("/api/counts?batch=v48g1").json()["mine_pending"] == 3
    q = rc.get("/api/cards?batch=v48g2&status=pending").json(); assert len([x for x in q if x.get("status") == "pending"]) == 1
