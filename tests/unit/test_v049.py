"""v0.4.9 regression tests (bot reports qa/*-v048.md). Made-up cards only."""
import io
import json
import re
from test_app import mk, sent


def _appr(ui, st, batch, **f):
    a = mk(ui, st, **f); ui.set_meta(a, batch=batch, batch_label=batch.upper()); st.update(a, status="approved", reviewer="RR"); return a


def _xlsx_text(b):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(b))
    return "\n".join(" ".join(str(c.value) for c in row if c.value is not None) for ws in wb.worksheets for row in ws.iter_rows())


def _surfaces(adm, batch, status):
    """Every place that shows a count, for one batch and filter."""
    q = f"status={status}&batch={batch}"
    out = {}
    bl = [b for b in adm.get("/api/batches").json() if b["batch"] == batch]
    out["batch_list"] = bl[0].get("ready") if bl and status == "approved_new" else None
    ck = adm.get(f"/api/export/check?{q}").json(); out["check"] = ck["ready"]["rows"]; out["check_rows"] = ck["rows"]
    sm = adm.get(f"/api/export/summary?{q}").json(); out["summary"] = sm["ready"]["rows"]
    v = adm.get(f"/api/export/xlsx?{q}&preview=1&view=1").json(); out["preview_rows"] = len(v["rows"]); out["preview_ready"] = v["ready"]["rows"]
    rd = adm.get(f"/api/export/xlsx?{q}&preview=1&companion=readme.txt").text
    m = re.search(r"Ready to import: (\d+) row", rd); out["readme"] = int(m.group(1)) if m else (0 if "Nothing new to import" in rd else None)
    sx = _xlsx_text(adm.get(f"/api/export/xlsx?{q}&preview=1&companion=summary.xlsx").content)
    m = re.search(r"Ready to import: (\d+) row", sx); out["summary_sheet"] = int(m.group(1)) if m else (0 if "Nothing new to import" in sx else None)
    dr = adm.get(f"/api/export/dryrun?{q}").text
    m = re.search(r"Ready to import: (\d+) row", dr); out["dryrun"] = int(m.group(1)) if m else (0 if "Nothing new to import" in dr else None)
    return out, ck, sm, rd


# ---- #1 one count everywhere
def test_v049_1_counts_agree_on_every_surface(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v49c", first_name="Ada", last_name="Sent", company="Sent Co", emails=["ada@v49c.example.com"])
    b = _appr(ui, st, "v49c", first_name="Ben", last_name="Fresh", company="Fresh Co", emails=["ben@v49c.example.com"])
    c = _appr(ui, st, "v49c", first_name="Cy", last_name="Edited", company="Edit Co", emails=["cy@v49c.example.com"])
    sent(ui, a, c); ui.set_meta(c, reexport=1)
    for status in ("approved_new", "approved"):
        out, ck, sm, rd = _surfaces(adm, "v49c", status)
        vals = {k: v for k, v in out.items() if v is not None}
        assert len(set(vals.values())) == 1 and set(vals.values()) == {1}, (status, out)   # only Ben goes in the file
    # status "approved": Ada (already sent) and Cy (edited after export) are left out, and the page says why in one line
    out, ck, sm, rd = _surfaces(adm, "v49c", "approved")
    assert ck["why"].startswith("Not in the file:") and "Left-out list" in ck["why"] and sm["why"] == ck["why"]


def test_v049_1_zero_rows_says_why(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    ids = [_appr(ui, st, "v49z", first_name=f"Zed{i}", last_name="Done", company=f"Z{i} Co", emails=[f"z{i}@v49z.example.com"]) for i in range(3)]
    sent(ui, *ids)
    ck = adm.get("/api/export/check?status=approved&batch=v49z").json()
    assert ck["ready"]["rows"] == 0 and ck["why"] == "3 cards already exported. Nothing new to import."
    ck2 = adm.get("/api/export/check?status=approved_new&batch=v49z").json()   # the default filter: no cards left, same line
    assert ck2["count"] == 0 and ck2["sent_before"] == 3 and ck2["why"] == ck["why"]
    out, *_ = _surfaces(adm, "v49z", "approved")
    assert set(v for v in out.values() if v is not None) == {0}, out
    from app import exporting
    assert exporting.why_line({"ready": {"rows": 0}}, 0) == "Nothing to export yet. Approve some cards first."


# ---- #2 CSV export is not "sent to MC Trade"
def test_v049_2_csv_export_never_counts_as_mc_trade(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v49csv", first_name="Cass", last_name="List", company="List Co", emails=["cass@v49csv.example.com"])
    r = adm.get("/api/export/csv?status=approved&batch=v49csv"); assert r.status_code == 200
    from app import exporting
    assert a not in exporting.sent_ids() and ui.meta(a).get("exported")
    assert any(h.get("action") == "exported_list" for h in ui.history(a))
    # an edit after a CSV-only export is NOT "edited after export"
    rr = adm.post(f"/api/cards/{a}", json={"fields": {"title": "Boss"}}).json()
    assert "Edited after export" not in (rr.get("notice") or "") and not ui.meta(a).get("reexport")
    ck = adm.get("/api/export/check?status=approved_new&batch=v49csv").json(); assert ck["ready"]["rows"] == 1
    # the IMPORT-ME export is the record; a CSV export afterwards never clears 'Edited after export'
    assert adm.get("/api/export/xlsx?status=approved_new&batch=v49csv").status_code == 200 and a in exporting.sent_ids()
    adm.post(f"/api/cards/{a}", json={"fields": {"title": "Chief"}}); assert ui.meta(a).get("reexport")
    adm.get("/api/export/csv?status=approved&batch=v49csv"); assert ui.meta(a).get("reexport")
    from app import audit_log
    assert "not sent to MC Trade" in audit_log._audit_text("exported_list", None, None, {"file": "x.csv", "kind": "csv"})


def test_v049_2_csv_check_column_carries_the_notes(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v49cc", first_name="Dot", last_name="Held", company="Held Co", emails=["dot@v49cc.example.com"])
    b = _appr(ui, st, "v49cc", first_name="Eve", last_name="New", company="New Co", emails=["eve@v49cc.example.com"])
    sent(ui, a)
    csv_ = adm.get("/api/export/xlsx?status=approved&batch=v49cc&preview=1&companion=extra.csv").content.decode("utf-8-sig")
    line = next(ln for ln in csv_.splitlines() if ln.startswith(f"{a},"))
    assert "Already exported" in line or "already exported" in line
    rd = adm.get("/api/export/xlsx?status=approved&batch=v49cc&preview=1&companion=readme.txt").text
    assert "see the CSV" not in rd or "Already exported" in csv_


# ---- #3 the same person across batches
def test_v049_3_company_suffix_and_plus_tag(env):
    from app import exporting, mctrade_sync as ms
    assert ms.norm_company("QA Skeptic Test Co.") == ms.norm_company("QA Skeptic Test Co") == ms.norm_company("qa skeptic test company")
    assert ms.norm_company("Acme, Inc.") == ms.norm_company("ACME LLC") == ms.norm_company("The Acme")
    assert ms.norm_email("Jo.Smith+gac@Example.com ") == "jo.smith@example.com"
    k1 = exporting.person_keys("Jo", "Smith", "Test Co.", ["jo+gac@ex49.example.com"])
    k2 = exporting.person_keys("jo", "SMITH", "Test Co", ["jo@ex49.example.com"])
    assert k1 & k2 and "e:jo@ex49.example.com" in k1 and any(k.startswith("n:") for k in k1 & k2)
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v49p1", first_name="Pat", last_name="Plus", company="Plus Test Co.", emails=["pat+gac@v49p.example.com"])
    sent(ui, a)
    b = _appr(ui, st, "v49p2", first_name="Pat", last_name="Plus", company="Plus Test Co", emails=["pat@v49p.example.com"])
    ck = adm.get("/api/export/check?status=approved_new&batch=v49p2").json()
    assert ck["ready"]["rows"] == 0 and ck["held"] == 1
    c = _appr(ui, st, "v49p3", first_name="Pat", last_name="Plus", company="Plus Test Company", emails=["other@v49q.example.com"])
    assert adm.get("/api/export/check?status=approved_new&batch=v49p3").json()["held"] == 1   # same name + same company


# ---- #4 privacy: another person's file name, batch or initials never shown
def test_v049_4_other_users_export_is_date_only(env):
    ui, st, adm, rev = env["ui"], env["store"], env["adm"], env["rev"]
    a = _appr(ui, st, "v49pv1", first_name="Priya", last_name="Vate", company="Vate Co", emails=["priya@v49pv.example.com"])
    sent(ui, a, user="someone-else", fname="GAC_IMPORT-ME_SECRET-BATCH_ZZ.xlsx", batch_label="SECRET-BATCH")
    b = mk(ui, st, first_name="Priya", last_name="Vate", company="Vate Co", emails=["priya@v49pv.example.com"]); ui.set_meta(b, batch="v49pv2", batch_label="V49PV2")
    st.update(b, status="approved", reviewer="RR")
    rd = rev.get("/api/export/xlsx?status=approved&batch=v49pv2&preview=1&companion=readme.txt").text
    sx = _xlsx_text(rev.get("/api/export/xlsx?status=approved&batch=v49pv2&preview=1&companion=summary.xlsx").content)
    sm = rev.get("/api/export/summary?status=approved&batch=v49pv2").json()
    for t in (rd, sx, json.dumps(sm)):
        assert "SECRET" not in t and "ZZ" not in t and "someone-else" not in t
    assert "in an earlier export (" in rd
    d = rev.get(f"/api/cards/{b}").json()
    assert "SECRET" not in json.dumps(d) and "someone-else" not in json.dumps(d)


# ---- #5 'Exported before' only when it applies, one line
def test_v049_5_exported_before_only_when_it_applies(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = _appr(ui, st, "v49x1", first_name="Xena", last_name="Once", company="Once Co", emails=["xena@v49x.example.com"])
    n = mk(ui, st, first_name="Nobody", last_name="Before", company="Fresh49 Co", emails=["nb@v49x.example.com"]); ui.set_meta(n, batch="v49x2")
    assert adm.get(f"/api/cards/{n}").json()["exported_elsewhere"] == []       # nothing applies: nothing shown
    adm.get("/api/export/csv?status=approved&batch=v49x1")                      # a contacts list is not an MC Trade export
    b = mk(ui, st, first_name="Xena", last_name="Once", company="Once Co", emails=["xena@v49x.example.com"]); ui.set_meta(b, batch="v49x2")
    assert adm.get(f"/api/cards/{b}").json()["exported_elsewhere"] == []
    sent(ui, a, user="adm", fname="GAC_IMPORT-ME_mine.xlsx")
    xe = adm.get(f"/api/cards/{b}").json()["exported_elsewhere"]
    assert len(xe) == 1 and xe[0]["last_file"] == "GAC_IMPORT-ME_mine.xlsx" and xe[0]["times"] == 1
    js = open("app/static/ui/views/review.js", encoding="utf-8").read()
    assert 'LEFT_OUT = "This person will be left out of the next file."' in js and "and ${" not in js.split("function xeLine")[1].split("}")[0]


# ---- #6 one wording for edited after export
def test_v049_6_one_wording(env):
    from app import ui, exporting, audit_log
    assert ui.EDITED_AFTER == "Edited after export, so it is not in the next file unless an admin includes it."
    assert exporting.CHANGED_WHY == "Edited after export, so it is not in this file unless an admin includes it."
    assert ui.EXPORTED_EDIT_NOTE == "Saved. " + ui.EDITED_AFTER
    assert audit_log._audit_text("reexport_needed", None, None, {}) == ui.EDITED_AFTER
    src = "".join(open(p, encoding="utf-8").read() for p in ("app/ui.py", "app/exporting.py", "app/audit_log.py", "app/static/ui/views/review.js", "app/static/ui/views/export.js"))
    assert "needs exporting again" not in src and "won't be exported again" not in src


# ---- #7 a typed batch name is used (New batch keeps it)
def test_v049_7_new_batch_keeps_typed_name():
    js = open("app/static/ui/views/upload.js", encoding="utf-8").read()
    clr = js.split('$("#clearBatch").onclick')[1].split("};")[0]
    assert 'if (!nm_ || (old_ && old_.startsWith(nm_))) $("#batchName").value = ""' in clr


def test_v049_7_server_uses_the_typed_name(env):
    rev = env["rev"]
    from PIL import Image
    b = io.BytesIO(); im = Image.new("RGB", (400, 250), "white")
    from PIL import ImageDraw
    d = ImageDraw.Draw(im); [d.text((20, 20 + 30 * i), f"Line {i} Example Co 555-0100", fill="black") for i in range(6)]; im.save(b, "PNG")
    r = rev.post("/api/jobs", data={"batch_label": "QA v049 coord"}, files={"image": ("a.png", b.getvalue(), "image/png")})
    assert r.status_code == 200, r.text
    assert r.status_code == 200 and r.json()["batch_label"] == "QA v049 coord"


# ---- #8 Needs admin wording follows the buttons that exist
def test_v049_8_needs_admin_messages():
    js = open("app/static/ui/views/review.js", encoding="utf-8").read()
    assert '" Fix it, or press Needs admin in the red box."' in js and '" Fix it, or press Confirm in the red box."' in js
    assert '$("#bNeedsAdmin")' in js and "[data-confirm]" in js
    assert '<span class="askadm">Ask an admin.</span>' not in js


# ---- #9 misreads: block only a clearly corrupted known word
def test_v049_9_misread_rules():
    from app import ui, plausibility as pl
    base = {"first_name": "Ann", "last_name": "Lee", "emails": ["a@b.example.com"], "phones": []}
    def blocked(fld, v): return ("name" if fld in ("first_name", "last_name") else fld) in ui.hard_errors(base | {fld: v})
    assert blocked("first_name", "Rnbert") and blocked("title", "Vice Presdient") and blocked("company", "Rnbert LLC")
    assert not blocked("last_name", "Rnjak") and not blocked("first_name", "Nrupen") and not blocked("company", "Managr Inc")
    assert not blocked("company", "Highlandisguare Cofee Co")
    hints = pl.check(base | {"company": "Highlandisguare Cofee Co"})
    assert any(h["code"] == "company_spelling" and "check spelling" in h["text"] for h in hints)
    assert pl.name_misread("Rnbert") and not pl.name_misread("Rnjak")


# ---- #10 small things
def test_v049_10_clip_at_word_boundary():
    from app import ui
    s = ui.clip("Brindlewood Ceramics and Pottery Supply Wholesale Rnbber", 40)
    assert s.endswith("…") and len(s) <= 41 and not s[:-1].endswith(" ") and "Brindlewood Ceramics and Pottery" in s
    assert ui.clip("short", 40) == "short"
    assert ui.clip("x" * 60, 40) == "x" * 40 + "…" or len(ui.clip("x" * 60, 40)) <= 41


def test_v049_10_no_stale_undo_after_fresh_sign_in(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    a = mk(ui, st, first_name="Una", last_name="Do", company="Undo Co", emails=["una@v49u.example.com"]); ui.set_meta(a, batch="v49u")
    assert adm.post(f"/api/cards/{a}", json={"fields": {"title": "Old edit"}}).status_code == 200
    assert adm.get("/api/undo?batch=v49u").json().get("id")
    from conftest import Client
    from app import main
    import time
    time.sleep(1.1)
    c2 = Client(main.app); assert c2.post("/api/auth/login", json={"username": "adm", "password": "admin-pass-123"}).status_code == 200
    assert c2.get("/api/undo?batch=v49u").json() == {}          # the old edit is not offered after a fresh sign-in


def test_v049_10_audit_lines_are_capped():
    from app import audit_log
    t = audit_log._audit_text("edited", "a" * 1000, "b" * 1000, {})
    assert len(t) < 700


def test_v049_10_copy_and_guides():
    up = open("app/static/ui/views/upload.js", encoding="utf-8").read()
    poll = up.split("async function poll")[1].split("\n}\n")[0] if "async function poll" in up else up
    assert "conn(false)" in poll
    import os
    V = open("VERSION").read().strip(); rd = open("README.md", encoding="utf-8").read(); assert V in rd
    src = open("app/exporting.py", encoding="utf-8").read() + open("app/static/ui/views/review.js", encoding="utf-8").read()
    for w in ("Affiliation code", "Columns not checked", "Test MC Trade list", "Test list"):
        assert w not in src, w
    ag = open("training/ADMIN-GUIDE.md", encoding="utf-8").read() if os.path.exists("training/ADMIN-GUIDE.md") else ""
    if ag: assert "real import template" in ag and "2 rows first" in ag


def test_v049_10_help_lengths():
    import re, os, pytest
    if not os.path.exists("training/REVIEWER-GUIDE.md"): pytest.skip("guides are built from training/ (private repo only)")
    def words(p):
        s = open(p, encoding="utf-8").read(); s = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", s); s = re.sub(r"\]\([^)]*\)", "]", s)
        return len(re.findall(r"[A-Za-z0-9][\w'’.-]*", s)), s
    n, s = words("training/REVIEWER-GUIDE.md"); assert n <= 450 and "demo video" in s and open("VERSION").read().strip() in s
    vis = len(re.findall(r"[A-Za-z0-9][\w'’.-]*", re.sub(r"<details.*?</details>", "", s, flags=re.S))); assert vis <= 400
    n, s = words("training/ADMIN-GUIDE.md"); assert n <= 800 and "2 rows first" in s
