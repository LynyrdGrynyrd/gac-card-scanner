"""Round-4 regression tests (reviewer-r4, skeptic-r4, newcomer-r4, mctrade-r1, manager-r1)."""
import io, csv
from test_app import mk
from openpyxl import load_workbook
from app import mctrade_export as M

def _approve(c, cid, **kw): return c.post(f"/api/cards/{cid}", json={"status": "approved", "crm_action": "create_new"} | kw)

def _note(st, cid, n):
    with st.db() as c: c.execute("UPDATE cards SET note=? WHERE id=?", (n, cid))

def test_build_dedupes_within_file(env):
    f = {"first_name": "Al", "last_name": "B", "emails": ["al@b.io"], "phones": [], "company": "B"}
    cards = [{"id": 1, "fields": f, "status": "approved"}, {"id": 2, "fields": dict(f), "status": "approved"}]
    res = M.build(cards, "X", "gac", env["ui"].hard_errors); rows, left = res["rows"], res["left"]
    assert len(rows) == 1 and left and "same person" in str(left).lower()

def test_safe_cell_keeps_phones(env):
    sc = env["ui"].safe_cell
    for v in ("+1 330-555-0181", "-5", "+44 (20) 7946 0958 x12", "330-555-0181", "plain"): assert sc(v) == v, v
    for v in ("=1+1", "+SUM(1)", "-2+3", "@SUM(1)", "\tx"): assert sc(v).startswith("'"), v
    assert sc("@cmd") == " @cmd"     # round 5: a bare @handle gets a leading space (text in Excel, no stray ')

def test_phone_and_address_helpers():
    assert M.phone("(330) 555-0181") == "330-555-0181" and M.phone("+1 330 555 0181", "e164") == "+13305550181"
    assert M.phone("330-555-0181 ext 22").endswith("x22")
    a = M.split_address("11000 Euclid Ave Cleveland OH 44106-7135")
    assert a["city"] == "Cleveland" and a["state"] == "OH" and a["zip"] == "44106-7135"
    assert M.clean('<b>Ann</b> <script>') == "Ann"

def test_api_status_approved_validated(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    cid = mk(ui, st, first_name="", last_name="Nobody", emails=["bad"])
    r = rev.post(f"/api/cards/{cid}", json={"status": "approved", "crm_action": "create_new"})
    assert r.status_code in (400, 409, 422) and ui.store.get(cid)["status"] == "pending"
    for path in ("/scan", "/card", "/sync/dry-run"): assert env["adm"].post(path, json={}).status_code in (404, 405, 410), path

def test_xlsx_template_dedupe_companions(env):
    ui, st, rev, adm = env["ui"], env["store"], env["rev"], env["adm"]
    a = mk(ui, st, first_name="Tia", last_name="Rowe", emails=["tia@rowe.io"], phones=[{"type": "work", "number": "+1 330-555-0181"}],
           address="11000 Euclid Ave Cleveland OH 44106-7135"); _note(st, a, "met at booth")
    b = mk(ui, st, first_name="Tia", last_name="Rowe", emails=["tia@rowe.io"])
    for c in (a, b): ui.set_meta(c, batch="r4x", batch_label="R4X")
    assert _approve(rev, a).status_code == 200
    rb = _approve(rev, b); assert rb.status_code == 409 and rb.json()["strong"]               # same person: must merge/reject
    assert rev.post(f"/api/cards/{b}", json={"status": "rejected"}).status_code == 200
    r = adm.get("/api/export/xlsx?status=approved&batch=r4x"); assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content)); assert wb.sheetnames == ["Sheet1"]; ws = wb.active
    hdr = [c.value or "" for c in ws[1]]; assert hdr == list(M.headers()[0])
    rows = [[c.value for c in x] for x in ws.iter_rows(min_row=2)]
    assert sum(1 for x in rows if "tia@rowe.io" in x) == 1                                   # one row per person
    row = rows[0]; g = lambda h: row[hdr.index(h)]
    assert g("IndFirstName") == "Tia" and g("IndCity") == "Cleveland" and g("IndWorkPhone") == "330-555-0181"
    assert g("AffCode1") == "CardScan-Pilot"
    eid = r.headers["X-Export-Id"]; assert set(r.headers["X-Companions"].split(",")) == {"extra.csv", "readme.txt"}
    ex = adm.get(f"/api/exports/{eid}/companion/extra.csv"); assert ex.text.startswith("\ufeff") and "met at booth" in ex.text
    assert env["rev2"].get(f"/api/exports/{eid}/companion/extra.csv").status_code == 403
    assert [e for e in adm.get("/api/exports").json() if e["id"] == int(eid)][0]["companions"]
    c = adm.get("/api/export/csv?status=approved&batch=r4x"); assert c.content.startswith(b"\xef\xbb\xbf") and "+1 330-555-0181" in c.text

def test_settings_affcode_phone(env):
    adm = env["adm"]
    assert adm.post("/api/admin/settings", json={"mctrade_affcode": "<x>"}).status_code in (400, 422)
    assert adm.post("/api/admin/settings", json={"phone_format": "nope"}).status_code in (400, 422)
    assert adm.post("/api/admin/settings", json={"mctrade_affcode": "GAC-2026", "phone_format": "e164"}).status_code == 200
    adm.post("/api/admin/settings", json={"mctrade_affcode": "CardScan-Pilot", "phone_format": "gac"})

def test_origin_scheme(env):
    r = env["adm"].post("/api/admin/settings", json={}, headers={"Origin": "https://testserver"})
    assert r.status_code == 403

def test_merge_keeps_note_and_purge_unmerges(env):
    ui, st, rev, adm = env["ui"], env["store"], env["rev"], env["adm"]
    t = mk(ui, st, first_name="Uma", last_name="Ng", emails=["uma@ng.io"]); _note(st, t, "target note")
    s = mk(ui, st, first_name="Uma", last_name="Ng", emails=["uma.priv@ng.io"]); _note(st, s, "dropped note")
    assert rev.post(f"/api/cards/{s}", json={"merge_into": t}).status_code == 200
    c = ui.store.get(t); assert "dropped note" in (c.get("note") or "") and "uma.priv@ng.io" in c["fields"]["emails"]
    r = adm.post(f"/api/admin/cards/{s}/delete", json={"confirm": s}); assert r.status_code == 200
    c = ui.store.get(t); assert "uma.priv@ng.io" not in c["fields"]["emails"] and "dropped note" not in (c.get("note") or "")

def test_batch_export_not_blocked_by_other_batch(env):
    ui, st, rev, rev2, adm = env["ui"], env["store"], env["rev"], env["rev2"], env["adm"]
    a = mk(ui, st, first_name="Vic", last_name="Oz", emails=["vic@oz.io"]); ui.set_meta(a, batch="bA")
    b = mk(ui, st, first_name="Vic", last_name="Oz", emails=["vic@oz.io"]); ui.set_meta(b, batch="bB")
    _approve(rev, a)
    assert rev.get("/api/export/check?status=approved&batch=bA").json()["duplicates"] == []
    assert rev.get("/api/export/check?status=rejected").status_code == 403
