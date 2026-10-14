"""Round-6 regression tests (qa/mctrade-r3.md)."""
import io, os, sys, csv, json, subprocess, threading
from test_app import mk
from openpyxl import load_workbook
from app import mctrade_export as M

def _card(i, decided=0, **f):
    base = {"first_name": "", "last_name": "", "company": "", "emails": [], "phones": []}; base.update(f)
    return {"id": i, "fields": base, "status": "approved", "decided": decided}
NOERR = lambda f: {}
OK = dict(first_name="James", last_name="Okafor", company="Cuyahoga Falls Tire Co.", emails=["jokafor@cftire.com"],
          phones=[{"type": "work", "number": "330-555-0177"}])
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# ---- workers: never on import, single instance
def test_import_starts_no_workers():
    code = ("import threading, app.main, app.ui as u; "
            "print(sorted(t.name for t in threading.enumerate() if t.name in ('ai-check','retention','recover')), u.WORKERS['started'])")
    out = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True, timeout=120)
    assert out.stdout.strip().endswith("[] False"), out.stdout + out.stderr

def test_worker_lock_single_instance(tmp_path):
    from app import ui
    p = str(tmp_path / "w.lock"); fd = ui._take_lock(p); assert fd is not None
    code = f"from app import ui; print(ui._take_lock({p!r}) is None)"
    out = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True, timeout=120)
    assert out.stdout.strip() == "True", out.stderr          # a second process can't take it
    os.close(fd)
    out = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True, timeout=120)
    assert out.stdout.strip() == "False"                       # released -> free again

def test_start_workers_refuses_when_locked(tmp_path, monkeypatch):
    from app import ui
    monkeypatch.setattr(ui, "LOCK_PATH", str(tmp_path / "x.lock")); monkeypatch.setattr(ui, "WORKERS", {"started": False, "lock_fd": None, "why": ""})
    holder = subprocess.Popen([sys.executable, "-c", f"from app import ui; import time; fd=ui._take_lock({str(tmp_path / 'x.lock')!r}); print('held', flush=True); time.sleep(30)"],
                              cwd=ROOT, stdout=subprocess.PIPE, text=True)
    try:
        assert holder.stdout.readline().strip() == "held"
        assert ui.start_workers() is False and "another process" in ui.WORKERS["why"]
    finally: holder.kill(); holder.wait()

# ---- export
def test_identical_kept_apart_exported_once_and_flagged():
    f = dict(first_name="Ann", last_name="Example", company="Sample U", emails=["ag@sample.example"], phones=[{"type": "work", "number": "216-555-0100"}])
    cs = [_card(76, 3, **f), _card(62, 2, **f), _card(34, 1, **f)]
    r = M.build(cs, "X", "gac", NOERR, kept=lambda a, b: True)
    assert len(r["rows"]) == 1 and r["rows"][0][0]["id"] == 76 and len(r["conflicts"]) == 2 and set(r["review"]) == {62, 34}

def test_different_ann_lees_same_company_stay_apart():
    a = _card(1, first_name="Ann", last_name="Lee", company="Acme", emails=["ann.lee@acme.com"])
    b = _card(2, first_name="Ann", last_name="Lee", company="Acme", emails=["alee2@acme.com"])
    assert len(M.build([a, b], "X", "gac", NOERR)["rows"]) == 2

def test_newest_decision_wins_merge():
    f = dict(first_name="Ann", last_name="Lee", company="Acme", emails=["ann@acme.com"])
    r = M.build([_card(1, 10, title="CEO", **f), _card(2, 20, title="Intern", **f)], "X", "gac", NOERR)
    assert r["rows"][0][1]["IndTitle"] == "Intern" and "IndTitle=CEO" in r["left"][0][1]

def test_updates_one_row_per_profile_and_no_create_for_same_person():
    cs = [_card(56, 4, **OK), _card(32, 3, **(OK | {"title": "Owner"})), _card(28, 2, **OK), _card(70, 1, **OK)]
    act = {56: "update_existing", 32: "update_existing", 28: "update_existing", 70: "create_new"}
    r = M.build(cs, "X", "gac", NOERR, crm_profile=lambda c: "1002", action=lambda c: act[c["id"]])
    assert len(r["updates"]) == 1 and r["updates"][0][1]["ProfileID"] == "1002" and r["updates"][0][1]["IndTitle"] == "Owner"
    assert r["rows"] == [] and 70 in r["review"]
    assert any(c["id"] == 70 and "NEEDS REVIEW" in why for c, why in r["left"])

def test_create_new_matching_crm_email_goes_to_review():
    # round 8 (mctrade-r5 F2): a card whose email matches an existing MC Trade profile is never filed as new: NEEDS REVIEW
    r = M.build([_card(67, **OK)], "X", "gac", NOERR, action=lambda c: "create_new", crm_exact=lambda c: "1002")
    assert r["rows"] == [] and r["review"] == [67] and r["decisions"][67][0] == "review" and "1002" in r["decisions"][67][1]

def test_strict_profile_id(env):
    ui = env["ui"]
    ex = [{"profile_id": "9", "first_name": "Robert", "last_name": "Hughes", "organization": "Hughes Law", "email": "rh@h.com", "phone": "330-555-0190"},
          {"profile_id": "8", "first_name": "Robbie", "last_name": "Hughes", "organization": "Other", "email": "x@o.com", "phone": "330-555-0111"}]
    c = lambda **f: {"id": 1, "fields": {"first_name": "Robert", "last_name": "Hughes", "company": "", "emails": [], "phones": []} | f}
    assert ui.strict_profile_id(c(emails=["rh@h.com"]), ex) == "9"                                         # exact email
    assert ui.strict_profile_id(c(phones=[{"type": "work", "number": "(330) 555-0190"}]), ex) == "9"       # exact name + phone
    assert ui.strict_profile_id(c(first_name="Robbie", phones=[{"type": "work", "number": "330-555-0190"}]), ex) is None  # namesake on a shared line
    assert ui.strict_profile_id(c(company="Hughes Law"), ex) is None                                        # name + org only

def test_update_without_certain_id_routed_to_review(env, monkeypatch):
    import os
    monkeypatch.setenv("MCTRADE_CRM_EXPORT", os.path.join(os.path.dirname(__file__), "fake_real_contacts.csv"))
    ui, st, rev, adm = env["ui"], env["store"], env["rev"], env["adm"]
    cid = mk(ui, st, first_name="Rob", last_name="Hughes", company="Hughes & Partners Law", emails=["rob.h@gmail.com"],
             phones=[{"type": "work", "number": "330-555-0190"}]); ui.set_meta(cid, batch="r6u", batch_label="R6U")
    assert rev.post(f"/api/cards/{cid}", json={"status": "approved", "crm_action": "update_existing", "confirm_cross_batch": True}).status_code == 200
    r = adm.get("/api/export/xlsx?status=approved&batch=r6u&preview=1"); assert r.status_code == 200
    assert "updates.xlsx" not in r.headers["X-Companions"]
    rd = adm.get("/api/export/xlsx?status=approved&batch=r6u&preview=1&companion=readme.txt").text
    assert "NEEDS REVIEW before importing" in rd and f"#{cid}" in rd

def test_email_strict_at_approval(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    for e in ("<img src=x>@acme.com", "=1+1@acme.com", "+SUM(1)@acme.com", "-x@acme.com", "a\u200b@acme.com"):
        cid = mk(ui, st, emails=[e]); r = rev.post(f"/api/cards/{cid}", json={"status": "approved", "crm_action": "create_new"})
        assert r.status_code == 422, (e, r.text)
    assert not ui.hard_errors({"first_name": "A", "emails": ["o'brien+tag@sub.acme.co.uk", "Ann@Acme.com ", "josé@exämple.de"]})

def test_content_checks_all_fields_and_depth():
    assert "last_name" in M.content_problems({"last_name": "&amp;amp;lt;img src=x&amp;amp;gt;"})
    assert "company" in M.content_problems({"company": "\uff1cimg src=x\uff1e"})
    assert "social" in M.content_problems({"social": ["<script>x</script>"]}) and not M.content_problems({"social": ["@gachamber"]})
    assert "other_address" in M.content_problems({"other_address": "<b>x</b>"})
    assert not M.content_problems({"note": "- call back Tuesday"})

def test_email_lowercased_in_export():
    r = M.build([_card(1, first_name="A", last_name="B", emails=["Ann@Acme.com "])], "X", "gac", NOERR)
    assert r["rows"][0][1]["IndEmail"] == "ann@acme.com"

# ---- dryrun
def test_dryrun_reads_updates_and_keeps_org_phone_blank(tmp_path):
    from openpyxl import Workbook
    hdr = M.FALLBACK
    def book(p, rows):
        wb = Workbook(); ws = wb.active; ws.append(hdr)
        for r in rows: ws.append([r.get(h, "") for h in hdr])
        wb.save(p)
    main = tmp_path / "GAC-MCTrade-x.xlsx"
    book(main, [{"IndFirstName": "Zeb", "IndLastName": "Quarrington", "OrgName": "Quarrington Widgets LLC", "IndEmail": "zeb@qw.example",
                 "IndWorkPhone": "330-555-0181"},
                {"IndFirstName": "Ann", "IndLastName": "Lawson", "OrgName": "L Co", "IndEmail": "ann@lco.example"}])
    book(tmp_path / "GAC-MCTrade-x-UPDATES.xlsx", [{"ProfileID": "1002", "IndFirstName": "James", "IndLastName": "Okafor",
          "OrgName": "Cuyahoga Falls Tire Co.", "IndEmail": "jokafor@cftire.com", "IndWorkPhone": "330-555-0177", "MemberSince": "1/1/1900"}])
    with open(tmp_path / "GAC-MCTrade-x-extra-columns.csv", "w", newline="") as fh:
        w = csv.writer(fh); w.writerow(["ScanCardID", "InImportFile", "FirstName", "LastName", "AllEmails", "MCTradeAction"])
        w.writerow([9, "no", "Ann", "Lawson", "ann.old@lco.example", ""])
        w.writerow([26, "yes", "Ann", "Lawson", "ann@lco.example", "create new"])
        w.writerow([31, "yes", "Zeb", "Quarrington", "zeb@qw.example", "create new"])
        w.writerow([56, "updates file", "James", "Okafor", "jokafor@cftire.com", "update existing"])
    out = subprocess.run([sys.executable, os.path.join(ROOT, "mctrade-dryrun", "dryrun.py"), "--export", str(main), "--out", str(tmp_path / "o"), "--run", "t"],
                         capture_output=True, text=True, timeout=120)
    assert out.returncode == 0, out.stderr
    assert "UPDATE 1" in out.stdout, out.stdout
    md = open(tmp_path / "o" / "t" / "match_report.md").read()
    assert "#26 | Ann Lawson" in md and "#31 | Zeb" in md and "#56 | James Okafor" in md
    ws = load_workbook(tmp_path / "o" / "t" / "upload" / "A_New_organizations_t.xlsx").active
    hdr2 = [c.value for c in ws[1]]; vals = [c.value for c in ws[2]]
    assert "330-555-0181" not in [str(v) for v in vals], dict(zip(hdr2, vals))

# ---- newcomer / reviewer / skeptic / manager r6
def test_save_with_wrong_card_or_stale_version_refused(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    cid = mk(ui, st, first_name="Ver", last_name="Sion", emails=["ver@sion.io"])
    v = rev.get(f"/api/cards/{cid}").json()["version"]
    r = rev.post(f"/api/cards/{cid}", json={"fields": {"title": "X"}, "card_id": cid + 1, "version": v}); assert r.status_code == 409 and r.json()["need"] == "wrong_card"
    assert rev.post(f"/api/cards/{cid}", json={"fields": {"title": "X"}, "card_id": cid, "version": v}).status_code == 200
    r = rev.post(f"/api/cards/{cid}", json={"fields": {"title": "Y"}, "card_id": cid, "version": v}); assert r.status_code == 409 and r.json()["need"] == "stale"
    assert st.get(cid)["fields"]["title"] == "X"

def test_client_via_ai_flag_ignored_and_no_phantom_note(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    cid = mk(ui, st, first_name="Silent", last_name="Edit")
    n0 = len(rev.get(f"/api/cards/{cid}").json()["history"])
    h = rev.post(f"/api/cards/{cid}", json={"fields": {"title": "Sneaky"}, "_via_ai": True}).json()["history"]
    assert len(h) == n0 + 1 and "Job title" in h[-1]["text"]
    h2 = rev.post(f"/api/cards/{cid}", json={"fields": {"title": "Sneaky2"}, "note": ""}).json()["history"]
    assert "Note" not in h2[-1]["text"]

def test_preview_scoped_for_reviewers_and_audited(env):
    ui, st, rev, adm = env["ui"], env["store"], env["rev"], env["adm"]
    assert rev.get("/api/export/xlsx?status=approved&preview=1").status_code == 403
    cid = mk(ui, st, first_name="Prev", last_name="Scope", emails=["prev@scope.io"]); ui.set_meta(cid, batch="r6p", batch_label="R6P")
    assert rev.post(f"/api/cards/{cid}", json={"status": "approved", "crm_action": "create_new"}).status_code == 200
    assert rev.get("/api/export/xlsx?status=approved&batch=r6p&preview=1").status_code == 200
    a = adm.get("/api/admin/audit?q=PREVIEW").json() if adm.get("/api/admin/audit?q=PREVIEW").status_code == 200 else None
    rows = ui.db().execute("SELECT event FROM auth_log WHERE event LIKE '%PREVIEW%' ORDER BY id DESC LIMIT 1").fetchall()
    assert rows and f"#{cid}" in rows[0][0]

def test_stale_ai_suggestion_refused(env, monkeypatch):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    from app import vlm, merge
    cid = mk(ui, st, first_name="Zebulonn"); p = os.path.join(ui.UP, "t_r6ai.jpg"); open(p, "wb").write(b"x")
    with st.db() as c: c.execute("UPDATE cards SET image='t_r6ai.jpg' WHERE id=?", (cid,))
    monkeypatch.setattr(vlm, "extract", lambda *a, **k: {})
    monkeypatch.setattr(merge, "merge", lambda f, v, raw: dict(f) | {"first_name": "Zebulon", "address": "2300 State Rd\nCuyahoga Falls, OH 44223"})
    ui._ai_run(cid, ["name"]); det = json.loads(ui.meta(cid)["ai_detail"])
    assert det["diff"]["address"]["to"] == "2300 State Rd, Cuyahoga Falls, OH 44223"
    rev.post(f"/api/cards/{cid}", json={"fields": {"first_name": "Zeb"}})
    r = rev.post(f"/api/cards/{cid}/ai-decide", json={"choices": {"first_name": "accept"}}); assert r.status_code == 409
    assert st.get(cid)["fields"]["first_name"] == "Zeb"; os.remove(p)

def test_same_photo_other_batch_warns_only(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    a = mk(ui, st, first_name="Pho", last_name="To", emails=["pho1@x.io"]); ui.set_meta(a, batch="ph1", sha256="r6sha")
    b = mk(ui, st, first_name="Other", last_name="Name", emails=["oth@y.io"], phones=[{"type": "work", "number": "216-555-0909"}]); ui.set_meta(b, batch="ph2", sha256="r6sha")
    x = [q for q in rev.get(f"/api/cards/{b}").json()["dupes"]["queue"] if q["card_id"] == a][0]
    assert "same photo" in x["why"] and not x["blocking"]

def test_img_long_and_control_names_404(env):
    rev = env["rev"]
    assert rev.get("/img/" + "a" * 255).status_code == 404 and rev.get("/img/a%01b.jpg").status_code == 404

def test_plausibility_false_positives_dropped():
    from app import plausibility as P
    codes = lambda f: {x["code"] for x in P.check({"first_name": "Ann", "last_name": "Lee", "phones": []} | f)}
    assert "email_tld" not in codes({"emails": ["a@x.om"]}) and "email_tld" not in codes({"emails": ["a@x.cm"]})
    assert "email_provider" not in codes({"emails": ["ann@ymail.com"]}) and "email_tld" in codes({"emails": ["a@x.con"]})
    assert "name_sanity" not in codes({"last_name": "Bank"})

def test_arithmetic_and_phone_formulas():
    assert M.looks_formula("+1 -1") and not M.looks_formula("+1 330-555-0181") and not M.looks_formula("-5")
    assert "phones" in M.content_problems({"phones": [{"type": "work", "number": "=1+2+3+4+5+6+7+8+9+0"}]})
    assert M.phone("=1+2+3+4+5+6+7+8+9+0") != "123-456-7890"

# ---- help-guide consistency (round 6)
def test_updates_name_has_no_new_and_inapp_dryrun(env):
    ui, st, rev, adm = env["ui"], env["store"], env["rev"], env["adm"]
    assert ui.comp_name("GAC-MCTrade-new-Expo-2026-10-13_2114.xlsx", "-UPDATES.xlsx") == "GAC-MCTrade-Expo-2026-10-13_2114-UPDATES.xlsx"
    assert ui.comp_name("GAC-MCTrade-new-Expo.xlsx", "-READ-ME.txt") == "GAC-MCTrade-new-Expo-READ-ME.txt"
    cid = mk(ui, st, first_name="Dry", last_name="Run", emails=["dry@run.io"]); ui.set_meta(cid, batch="r6d", batch_label="R6D")
    assert rev.post(f"/api/cards/{cid}", json={"status": "approved", "crm_action": "create_new", "confirm_cross_batch": True}).status_code == 200
    assert rev.get("/api/export/dryrun?status=approved").status_code == 403          # reviewers: one batch
    r = rev.get("/api/export/dryrun?status=approved&batch=r6d"); assert r.status_code == 200 and "match report" in r.text and "Dry Run" in r.text

def test_template_upload_checked(env, tmp_path, monkeypatch):
    from openpyxl import Workbook
    adm, rev = env["adm"], env["rev"]
    monkeypatch.setattr(M, "TEMPLATE_XLSX", str(tmp_path / "mctrade_template.xlsx"))
    wb = Workbook(); wb.active.title = "Instructions"; wb.active.append(["Read me"]); ws = wb.create_sheet("Import"); ws.append(["Profile ID", "IndFirstName", "IndLastName", "IndEmail ", "Category", "MyCustom"])
    p = tmp_path / "t.xlsx"; wb.save(p)
    assert rev.post("/api/admin/mctrade-template", files={"image": ("t.xlsx", open(p, "rb"))}).status_code in (401, 403)
    assert adm.post("/api/admin/mctrade-template", files={"image": ("x.txt", b"hello")}).status_code == 422
    r = adm.post("/api/admin/mctrade-template", files={"image": ("t.xlsx", open(p, "rb"))}); assert r.status_code == 200, r.text
    assert r.json()["columns"][0] == "Profile ID" and "MyCustom" in r.json()["unknown"]       # instructions sheet skipped
