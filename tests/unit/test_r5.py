"""Round-5 regression tests (mctrade-r2, skeptic-r5, newcomer-r5, reviewer-r5, manager-r2)."""
import io, csv, json
from test_app import mk
from openpyxl import load_workbook
from app import mctrade_export as M

def _approve(c, cid, **kw): return c.post(f"/api/cards/{cid}", json={"status": "approved", "crm_action": "create_new"} | kw)
def _card(i, **f):
    base = {"first_name": "", "last_name": "", "company": "", "emails": [], "phones": []}; base.update(f)
    return {"id": i, "fields": base, "status": "approved"}
NOERR = lambda f: {}

# ---- MC Trade export
def test_clean_decodes_entities_then_strips():
    assert M.clean("Ann &lt;img src=x onerror=alert(1)&gt; Lee") == "Ann Lee"
    assert M.clean("&amp;lt;script&amp;gt;x") == "x" and M.clean("5 < 6") == "5 < 6" and M.clean("Strand<img src=x>") == "Strand"
    assert M.content_problems({"last_name": "Lee&lt;script&gt;"}) and M.content_problems({"title": "=1+1"})
    assert not M.content_problems({"title": "-5", "company": "+1 330-555-0181"})

def test_dedup_shared_role_email_and_unicode_names():
    a = _card(1, first_name="Ann", last_name="Lee", emails=["info@acme.com"], company="Acme")
    b = _card(2, first_name="Bob", last_name="Ray", emails=["info@acme.com"], company="Acme")
    c = _card(3, first_name="张", last_name="伟", company="华为", phones=[{"type": "work", "number": "330-555-0101"}])
    d = _card(4, first_name="李", last_name="娜", company="华为", phones=[{"type": "work", "number": "330-555-0102"}])
    e = _card(5, first_name="José", last_name="Núñez", company="X", emails=["j@x.io"])
    g = _card(6, first_name="Jos", last_name="Nez", company="X", emails=["jn@x.io"])
    r = M.build([a, b, c, d, e, g], "X", "gac", NOERR)
    assert len(r["rows"]) == 6, r["left"]

def test_dedup_respects_keep_apart_and_merges_details():
    f = dict(first_name="Ann", last_name="Example", company="Sample U", emails=["ag@sample.example"], phones=[{"type": "work", "number": "216-555-0100"}])
    a, b = _card(62, **f), _card(34, **(f | {"phones": [{"type": "work", "number": "216-555-0100"}, {"type": "mobile", "number": "216-555-0188"}, {"type": "fax", "number": "216-555-0199"}]}))
    r = M.build([a, b], "X", "gac", NOERR)
    assert len(r["rows"]) == 1 and r["rows"][0][1]["IndCellPhone"] == "216-555-0188" and r["rows"][0][1]["IndFax"] == "216-555-0199"
    assert "same person as card #62" in r["left"][0][1]
    r = M.build([a, b], "X", "gac", NOERR, kept=lambda x, y: {x, y} == {62, 34})
    assert len(r["rows"]) == 2

def test_update_existing_goes_to_updates_file(env):
    ui, st, rev, adm = env["ui"], env["store"], env["rev"], env["adm"]
    cid = mk(ui, st, first_name="Robert", last_name="Hughes", company="Hughes & Partners Law", emails=["rhughes@hugheslaw.com"],
             phones=[{"type": "work", "number": "330-555-0190"}]); ui.set_meta(cid, batch="r5u", batch_label="R5U")
    assert _approve(rev, cid, crm_action="update_existing").status_code == 200, _approve(rev, cid, crm_action="update_existing").text
    r = adm.get("/api/export/xlsx?status=approved&batch=r5u"); assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active; assert ws.max_row == 1           # nothing in the NEW-profiles file
    assert "updates.xlsx" in r.headers["X-Companions"]
    u = adm.get(f"/api/exports/{r.headers['X-Export-Id']}/companion/updates.xlsx"); assert u.status_code == 200
    uw = load_workbook(io.BytesIO(u.content)).active; hdr = [c.value for c in uw[1]]; row = [c.value for c in uw[2]]
    assert row[hdr.index("ProfileID")] == "1004" and row[hdr.index("MemberSince")] == "1/1/1900"

def test_formula_and_html_refused_at_approval(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    for f in ({"title": '=HYPERLINK("http://e","x")'}, {"company": "+SUM(1+1)"}, {"website": "@SUM(1)"}, {"last_name": "Lee&lt;img src=x&gt;"}):
        cid = mk(ui, st, **f); r = _approve(rev, cid); assert r.status_code == 422, (f, r.text)

def test_preview_export_marks_nothing(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    cid = mk(ui, st, first_name="Pre", last_name="View", emails=["pre@view.io"]); ui.set_meta(cid, batch="r5p", batch_label="R5P")
    assert _approve(rev, cid).status_code == 200
    n = len(rev.get("/api/exports").json())
    r = rev.get("/api/export/xlsx?status=approved_new&batch=r5p&preview=1"); assert r.status_code == 200 and r.headers["X-Preview"] == "1"
    assert not ui.meta(cid).get("exported") and len(rev.get("/api/exports").json()) == n

def test_csv_guard_dashed_phone_and_handles(env):
    sc = env["ui"].safe_cell
    assert sc("+1-330-555-0182") == "'+1-330-555-0182" and sc("+1-1").startswith("'")
    assert sc("+1 330-555-0181") == "+1 330-555-0181" and sc("-5") == "-5" and sc("@fphealthakron") == " @fphealthakron"

def test_short_filenames(env):
    fn = env["ui"]._fname("contacts-for-MC-Trade", "approved_new", "xlsx", "R5 QA newcomer Oct 13 · Oct 13, 8:22 PM · QN")
    assert len(fn) < 60 and fn.count("Oct-13") == 1, fn

# ---- skeptic
def test_reapprove_keeps_override(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    cid = mk(ui, st, first_name="Zeb", last_name="Quarrington", emails=["zeb@quarrington"])
    assert _approve(rev, cid, override_reason="QA: domain confirmed").status_code == 200
    r = rev.post(f"/api/cards/{cid}", json={"status": "approved"}); assert r.status_code == 200
    assert r.json()["override"]["reason"] == "QA: domain confirmed"

def test_img_nul_is_404(env):
    assert env["rev"].get("/img/%00").status_code == 404

def test_merge_of_approved_needs_confirm_and_note_dropped(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    a = mk(ui, st, first_name="Mo", last_name="One", emails=["mo@one.io"]); b = mk(ui, st, first_name="Pell", last_name="Two", emails=["p@two.io"])
    assert _approve(rev, a).status_code == 200
    r = rev.post(f"/api/cards/{a}", json={"merge_into": b}); assert r.status_code == 409 and r.json()["need"] == "confirm_redecide"
    assert rev.post(f"/api/cards/{a}", json={"merge_into": b, "confirm_redecide": True}).status_code == 200
    assert "Duplicate of" in st.get(a)["note"]
    r = rev.post(f"/api/cards/{a}", json={"status": "approved", "confirm_redecide": True, "keep_apart": "all"}); assert r.status_code == 200, r.text
    assert "Duplicate of" not in (st.get(a)["note"] or "")

def test_privacy_delete_purges_target_snapshots_and_exports(env):
    ui, st, rev, adm = env["ui"], env["store"], env["rev"], env["adm"]
    a = mk(ui, st, first_name="Del", last_name="Me", emails=["delme@x.io"], phones=[{"type": "mobile", "number": "330-555-7777"}])
    b = mk(ui, st, first_name="Keep", last_name="Target", emails=["keep@t.io"]); ui.set_meta(b, batch="r5d", batch_label="R5D")
    assert rev.post(f"/api/cards/{a}", json={"merge_into": b}).status_code == 200
    assert rev.post(f"/api/cards/{b}", json={"fields": {"title": "Boss"}}).status_code == 200       # snapshot of b holds delme@x.io
    assert _approve(rev, b).status_code == 200
    assert rev.get("/api/export/csv?status=approved&batch=r5d").status_code == 200
    assert adm.post(f"/api/admin/cards/{a}/delete", json={"confirm": a}).status_code == 200
    with ui.db() as x:
        for s in x.execute("SELECT snapshot FROM actions").fetchall(): assert "delme@x.io" not in (s[0] or "")
        for e in x.execute("SELECT stored, card_ids FROM exports").fetchall():
            if b in json.loads(e[1] or "[]"): assert e[0] is None
    assert "delme@x.io" not in json.dumps(st.get(b)["fields"])

# ---- newcomer / reviewer / manager
def test_cross_batch_lookalike_does_not_block_approval(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    a = mk(ui, st, first_name="Cross", last_name="Batch", emails=["cb@x.io"]); ui.set_meta(a, batch="other", batch_label="Other")
    assert _approve(rev, a).status_code == 200
    b = mk(ui, st, first_name="Cross", last_name="Batch", emails=["cb@x.io"]); ui.set_meta(b, batch="mine", batch_label="Mine")
    r = rev.post("/api/cards-bulk", json={"ids": [b], "status": "approved"}); assert b in r.json()["updated"], r.json()
    c = mk(ui, st, first_name="Cross", last_name="Batch", emails=["cb@x.io"]); ui.set_meta(c, batch="other", batch_label="Other")
    assert _approve(rev, c).status_code == 409                                                   # same batch still blocks

def test_note_edit_logged_and_bulk_label(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    cid = mk(ui, st, first_name="No", last_name="Te")
    assert rev.post(f"/api/cards/{cid}", json={"note": "called back"}).status_code == 200
    assert any(h["action"] == "edited" and "called back" in json.dumps(h) for h in ui.history(cid))
    rev.post("/api/cards-bulk", json={"ids": [cid], "status": "rejected"})
    assert "Bulk reject 1" in rev.get("/api/undo").json().get("label", "")

def test_back_of_other_person_asks(env, monkeypatch):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    cid = mk(ui, st, first_name="Priya", last_name="Raman", emails=["priya@r.io"])
    monkeypatch.setattr(ui, "_check_image", lambda d: None)
    monkeypatch.setattr(ui, "scan_limited", lambda *a, **k: {"text": "Robert Hughes", "fields": {"first_name": "Robert", "last_name": "Hughes", "emails": ["rh@h.io"], "phones": []}})
    r = rev.post(f"/api/cards/{cid}/back", files={"image": ("b.jpg", b"x" * 10, "image/jpeg")})
    assert r.status_code == 409 and r.json()["need"] == "confirm_other_person" and "rh@h.io" not in json.dumps(st.get(cid)["fields"])
    r = rev.post(f"/api/cards/{cid}/back", files={"image": ("b.jpg", b"x" * 10, "image/jpeg")}, data={"confirm_other_person": "1"})
    assert r.status_code == 200

def test_reset_password_lifts_lock_and_retention_min(env):
    adm = env["adm"]; from app import auth
    uid = [u for u in adm.get("/api/admin/users").json() if u["username"] == "rev2"][0]["id"]
    import time; auth._fails["rev2"] = [time.time()] * 6
    assert adm.post(f"/api/admin/users/{uid}", json={"password": "New-pass-98765x"}).status_code == 200
    assert "rev2" not in auth._fails
    assert adm.post("/api/admin/settings", json={"image_retention_days": "0"}).status_code == 422
    assert adm.post("/api/batches/archive", json={"batch": "__all__", "confirm": "clear"}).status_code == 422

def test_health_reports_ai_checker(env):
    j = env["rev"].get("/health").json(); assert j["ok"] and j["ai_checker"] in ("online", "offline")
