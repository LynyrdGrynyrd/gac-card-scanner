"""Round-3 tester findings (qa/*-r3.md). Run: .venv/bin/python -m pytest tests/unit"""
import io, os, json, csv
from test_app import mk

def _approve(c, cid, **kw): return c.post(f"/api/cards/{cid}", json={"status": "approved", "crm_action": "create_new"} | kw)

def test_formula_injection_csv_and_xlsx(env):
    ui, st, adm = env["ui"], env["store"], env["adm"]
    cid = mk(ui, st, first_name="Oph", last_name="Strand", title='=HYPERLINK("http://evil.example","x")', company="+SUM(1+1)",
             website="@SUM(1)", address="-2+3+cmd|' /C calc'!A0", emails=["o@strand.io", "o2@strand.io"])
    ui.set_meta(cid, batch="fx", batch_label="FX Event")
    r5 = _approve(env["rev"], cid); assert r5.status_code == 422 and "formula" in r5.text   # round 5: formula text can't be approved
    st.update(cid, status="approved", reviewer="OLD", note="")                             # ... but older approved data may hold it
    r = adm.get("/api/export/csv?status=approved&batch=fx"); assert r.status_code == 200
    assert "FX-Event" in r.headers["content-disposition"]                                  # batch in the file name
    row = [x for x in csv.reader(io.StringIO(r.text)) if x and x[0] == str(cid)][0]
    assert not [v for v in row if v[:1] in ("=", "+", "@", "-", "\t", "\r")], row
    assert "'=HYPERLINK" in r.text and "o2@strand.io" in row                                  # email2 column
    r = adm.get("/api/export/xlsx?status=approved&batch=fx"); assert r.status_code == 200
    assert r.json()["nothing"] and "1 to fix first" in r.json()["ready"]["text"]          # v0.4.8 #1: no empty IMPORT-ME file
    ok = mk(ui, st, first_name="Clean", last_name="Row", company="Fine Co", emails=["clean@row.io"]); ui.set_meta(ok, batch="fx", batch_label="FX Event")
    st.update(ok, status="approved", reviewer="OLD", note="")
    r = adm.get("/api/export/xlsx?status=approved&batch=fx"); assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content)); assert wb.sheetnames == ["Sheet1"]; ws = wb.active   # template only (round 4)
    for row in ws.iter_rows(min_row=2):
        for c in row: assert c.data_type != "f", (c.coordinate, c.value)
    eid = r.headers["X-Export-Id"]; ex = adm.get(f"/api/exports/{eid}/companion/extra.csv")
    assert ex.status_code == 200 and "o2@strand.io" in ex.text                             # second email -> extra CSV

def test_reviewer_export_filters_and_downloads(env):
    rev, rev2, adm = env["rev"], env["rev2"], env["adm"]
    for s in ("rejected", "all", "pending", "anything"): assert rev.get(f"/api/export/csv?status={s}").status_code == 403, s
    assert rev.get("/api/export/check?status=rejected").json()["admin_only"]
    ui, st = env["ui"], env["store"]; cid = mk(ui, st, first_name="Dl", last_name="Own", emails=["dl@own.io"]); ui.set_meta(cid, batch="dl")
    _approve(rev, cid); r = rev.get("/api/export/csv?status=approved&batch=dl"); assert r.status_code == 200
    eid = [e for e in rev.get("/api/exports").json() if e["username"] == "rev"][0]["id"]
    assert rev.get(f"/api/exports/{eid}/download").status_code == 200
    assert rev2.get(f"/api/exports/{eid}/download").status_code == 403
    assert adm.get(f"/api/exports/{eid}/download").status_code == 200
    assert not [e for e in rev2.get("/api/exports").json() if e["id"] == eid]   # v0.4.3 #8: not listed for other reviewers

def test_edit_or_merge_into_approved_rechecks(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    a = mk(ui, st, first_name="Zeb", last_name="Q", emails=["zeb@q.io"]); assert _approve(rev, a).status_code == 200
    r = rev.post(f"/api/cards/{a}", json={"fields": {"emails": ["garbage-not-email"], "phones": [{"type": "work", "number": "12"}]}})
    assert r.status_code == 200 and r.json()["status"] == "pending" and "moved back to pending" in r.json()["notice"]
    assert any(h["action"] == "back_to_pending" for h in r.json()["history"])
    t = mk(ui, st, first_name="Mar", last_name="Del", emails=["mar@del.io"], company="Del Co"); assert _approve(rev, t).status_code == 200
    s = mk(ui, st, first_name="Mar", last_name="Del", emails=["no-at-sign", "mar.home@del.io"], phones=[{"type": "work", "number": "9"}], company="Del Co")
    ui.set_meta(s, back_image="back_s.jpg")
    assert rev.post(f"/api/cards/{s}", json={"merge_into": t}).status_code == 200
    # round 12 (coordinator r4 #3): the invalid email/phone are no longer merged in, so the approved card has no new problem
    # and stays approved; the valid extra email and the back photo are carried over
    assert st.get(t)["status"] == "approved" and ui.meta(t)["back_image"] == "back_s.jpg"
    assert st.get(t)["fields"]["emails"] == ["mar@del.io", "mar.home@del.io"]

def test_first_name_hard_and_all_errors_at_once(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    cid = mk(ui, st, first_name="", last_name="Lee", emails=["bad@@"], phones=[{"type": "work", "number": "12"}])
    r = _approve(rev, cid, override_reason="sure"); j = r.json()
    assert r.status_code == 422 and j["need"] == "fix" and {"name", "emails", "phones"} <= set(j["errors"])
    cid2 = mk(ui, st, first_name="", last_name="Lee", emails=["ok@lee.io"])
    r = _approve(rev, cid2, override_reason="sure"); assert r.status_code == 422 and "name" in r.json()["errors"]

def test_keep_both_per_match_and_true_duplicates(env):
    ui, st, rev = env["ui"], env["store"], env["rev"]
    PH = [{"type": "work", "number": "330-555-0101"}]   # shared line (mk gives unique phones since round 7)
    a = mk(ui, st, first_name="James", last_name="Okafor", emails=["james@ok.io"], company="Ok Ltd", phones=PH); ui.set_meta(a, batch="kb")
    b = mk(ui, st, first_name="James", last_name="Okafor", emails=["james@ok.io"], company="Ok Ltd", phones=PH); ui.set_meta(b, batch="kb")    # same batch, same email
    c = mk(ui, st, first_name="James", last_name="Okafor", emails=["jo@other.org"], company="Ok Ltd", phones=PH); ui.set_meta(c, batch="other")
    d = env["adm"].get(f"/api/cards/{a}").json()["dupes"]["queue"]   # v0.4.7 #1: a reviewer sees only this batch's card in full
    assert {x["card_id"]: x["strong"] for x in d if x["card_id"] in (b, c)} == {b: True, c: False}
    d = rev.get(f"/api/cards/{a}").json()["dupes"]["queue"]
    assert {x["card_id"]: x["strong"] for x in d if x.get("card_id") in (b, c)} == {b: True}
    r = rev.post(f"/api/cards/{a}", json={"keep_apart": [b]}); assert r.status_code == 409 and r.json()["need"] == "true_duplicate"
    r = rev.post(f"/api/cards/{a}", json={"dup_resolution": "keep_both"}); assert r.status_code == 409          # legacy all-at-once too
    r = rev.post(f"/api/cards/{a}", json={"keep_apart": [c]}); assert r.status_code == 200
    q = {x["card_id"]: x["kept_apart"] for x in env["adm"].get(f"/api/cards/{a}").json()["dupes"]["queue"]}; assert q[c] and not q[b]
    assert r.json()["dupes"]["unresolved_queue"]

    for x in (a, b): rev.post(f"/api/cards/{x}", json={"status": "approved", "crm_action": "create_new"})
    assert rev.get("/api/export/csv?status=approved&batch=kb").status_code == 409

def test_ai_never_writes_and_priority(env, monkeypatch):
    ui, st = env["ui"], env["store"]
    from app import vlm, merge
    cid = mk(ui, st, title="Mgr", department="")
    p = os.path.join(ui.UP, "t_ai3.jpg"); open(p, "wb").write(b"x")
    with st.db() as c: c.execute("UPDATE cards SET image='t_ai3.jpg' WHERE id=?", (cid,))
    monkeypatch.setattr(vlm, "extract", lambda *a, **k: {})
    monkeypatch.setattr(merge, "merge", lambda f, v, raw: dict(f) | {"title": "Manager", "department": "Ops"})
    env["rev"].post(f"/api/cards/{cid}", json={"fields": {"department": "Tax"}})          # person edits during the run
    f0 = st.get(cid)["fields"]; ui._ai_run(cid, ["title"])
    assert st.get(cid)["fields"] == f0                                                      # pending + unedited fields: still not written
    det = json.loads(ui.meta(cid)["ai_detail"]); assert det["diff"]["department"]["edited_by_person"] and ui.meta(cid)["ai_status"] == "suggested"
    h = env["rev"].get(f"/api/cards/{cid}").json()["history"]
    assert any("suggestions offered for" in x["text"] for x in h)
    r = env["rev"].post(f"/api/cards/{cid}/ai-decide", json={"choices": {"title": "accept"}, "record_only": True})
    assert st.get(cid)["fields"]["title"] == "Mgr" and "AI suggestion: Job title: used the AI value" in r.json()["history"][-1]["text"]
    r = env["rev"].post(f"/api/cards/{cid}/ai-decide", json={"choices": {"department": "reject"}})
    assert "department" not in r.json()["low"]
    # priority: a manual request jumps ahead of background ones
    ui.AI_WAIT.clear(); ui.AI_WAIT.update({901: (1, 1), 902: (1, 2)}); ui.AI_WAIT[903] = (0, 3)
    assert ui.ai_progress(903)["ahead"] == 0 and ui.ai_progress(902)["ahead"] == 2
    ui.AI_WAIT.clear(); os.remove(p)

def test_sync_admin_only_and_already_scanned(env):
    from PIL import Image
    rev = env["rev"]
    assert rev.get("/sync/dry-run/1").status_code == 403
    b = io.BytesIO(); Image.effect_noise((900, 500), 60).convert("RGB").save(b, "JPEG")
    from app import ui
    ui.set_meta(999999, sha256=__import__("hashlib").sha256(b.getvalue()).hexdigest())  # orphan meta has no card: ignored
    cid = mk(ui, env["store"], by=""); ui.set_meta(cid, sha256=__import__("hashlib").sha256(b.getvalue()).hexdigest(), batch_label="Expo", batch="b-expo")
    ui.log(cid, "QN", "scanned", None, "pending", {"file": "x.jpg"})
    r = rev.post("/api/jobs", data={"batch": "b-expo"}, files={"image": ("x.jpg", b.getvalue(), "image/jpeg")}); assert r.status_code == 409
    assert "by QN on" in r.json()["error"] and "Expo" in r.json()["error"]

def test_hard_delete_scrubs_merge_history_and_jobs(env):
    ui, st, adm, rev = env["ui"], env["store"], env["adm"], env["rev"]
    t = mk(ui, st, first_name="Ign", last_name="Pell", emails=["ign@pell.io"])
    s = mk(ui, st, first_name="Ign", last_name="Pell", emails=["secret@pell.io"])
    rev.post(f"/api/cards/{s}", json={"merge_into": t})
    json.dump({"id": "jdel", "card_id": s, "file": "secret-name.jpg"}, open(os.path.join(ui.DONE, "jdel.json"), "w"))
    ui.JOBS["jdel"] = {"id": "jdel", "state": "done", "card": {"id": s}}
    assert adm.post(f"/api/admin/cards/{s}/delete", json={"confirm": s}).status_code == 200
    with st.db() as c:
        assert not [r for r in c.execute("SELECT detail FROM card_history WHERE card_id=?", (t,)) if r[0] and "secret@pell.io" in r[0]]
    assert not os.path.exists(os.path.join(ui.DONE, "jdel.json")) and "jdel" not in ui.JOBS

def test_html_no_store(env):
    r = env["rev"].get("/ui"); assert r.headers["cache-control"] == "no-store"
