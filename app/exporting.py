"""Exports: CSV, vCard and the MC Trade Excel, sign-off versions, export history and re-download. Split out of app/ui.py (audit 2026-10 P2)."""
import io, os, re, csv, json, time, uuid, hashlib, datetime, threading
from fastapi import HTTPException, Request
from fastapi.responses import FileResponse, Response, JSONResponse
from . import store, export, mctrade_sync, auth, crm_source
from . import ui as U

# ---------------------------------------------------------------- exports (reviewer, duplicate guard, history)
XCOLS = ["IndFirstName", "IndLastName", "IndTitle", "OrgName", "IndEmail", "IndEmail2", "IndPhone", "IndCellPhone", "IndFax",
         "WebAddress", "Address", "AffCode1"]
MORE = ["Credentials", "Department", "OtherAddress", "AllEmails", "AllPhones", "Note"]
AUDIT = ["ScanCardID", "ReviewedBy", "DecidedAt", "Batch", "MCTradeAction", "Check"]
FILTERS = {"approved_new": "approved, not yet exported", "approved": "all approved", "all": "all cards", "pending": "pending only", "rejected": "rejected only"}

def _export_set(status, batch=""):
    cs, metas = U._active()
    mm = lambda c: metas.get(c["id"]) or {}
    # v0.4.9 #2: 'not yet exported' = not yet in an MC Trade import file (a contacts list doesn't count), or edited since
    if status == "approved_new": sent = sent_ids(); sel = [c for c in cs if c["status"] == "approved" and (c["id"] not in sent or mm(c).get("reexport"))]
    else: sel = [c for c in cs if status == "all" or c["status"] == status]
    if batch: sel = [c for c in sel if (mm(c).get("batch") or "__none__") == batch]
    probs = []
    if batch:   # a batch export is only blocked by look-alikes inside the same batch; other people's batches don't block it
        inb = {c["id"] for c in cs if (mm(c).get("batch") or "__none__") == batch}
    for c in sel:   # all batches: look-alikes anywhere in the active queue (pending ones too)
        d = U.dupes(c, cs, metas)
        if batch: d["open"] = [x for x in d["open"] if x["card_id"] in inb]
        # round 15 (reviewer r14): rejected look-alikes stay as review hints but never block an export
        clash = [x["card_id"] for x in d["open"] if (store.get(x["card_id"]) or {}).get("status") != "rejected"]
        if clash: probs.append({"id": c["id"], "with": clash, "strong": [x["card_id"] for x in d["open"] if x.get("strong") and (store.get(x["card_id"]) or {}).get("status") != "rejected"],
                                "batch": mm(c).get("batch_label") or mm(c).get("batch") or ""})
    return sel, metas, probs

ADMIN_ONLY_FILTERS = ("all", "pending")
REVIEWER_FILTERS = ("approved_new", "approved")      # anything else (rejected, any free text) is admin-only
def _filter_refused(status, user):
    if user["role"] == "admin": return status not in FILTERS and status != "rejected"
    return status not in REVIEWER_FILTERS

def batch_owner(batch):
    """Initials of whoever scanned the batch's first card ('' if unknown)."""
    r = U.db().execute("""SELECT h.reviewer FROM card_meta m JOIN card_history h ON h.card_id=m.card_id AND h.action='scanned'
        WHERE m.batch=? ORDER BY h.id LIMIT 1""", (batch,)).fetchone()
    return str(r[0] or "") if r else ""
BATCH_MSG = "You can export only your own batch. Ask an admin to export other batches or all batches."
def _batch_refused(batch, user):
    """v0.4.2 #1 (mctrade-admin M1): enforced here, not only in the page. A reviewer exports only a batch they started;
    another person's batch, cards with no batch and all batches are admin-only."""
    if user.get("role") == "admin": return False
    from . import demo
    if demo.enabled(): return False   # public demo: every export is a preview of fictional cards
    if not batch or batch == "__none__": return True
    oid = U.batch_owner_id(batch)   # v0.4.3 #8: by user id
    return oid != user.get("id") if oid is not None else batch_owner(batch) != user.get("initials")

PLAINNUM = re.compile(r"[+-]\d+(?:\.\d+)?"); HANDLE = re.compile(r"@[A-Za-z0-9_.]{1,40}(?:\s+@[A-Za-z0-9_.]{1,40})*")

NUMBERISH = re.compile(r"[+-][\d\s().-]*\d[\d\s().-]*(?:\s*(?:x|ext\.?|extension|#)\s*\d+)?", re.I)
def safe_cell(v):
    """CSV formula-injection guard. A leading ' is added ONLY when a spreadsheet would evaluate the text: it starts with
    = or @ or TAB/CR, or with + / - followed by something that is not a plain phone number or number
    ('+1 330-555-0181' and '-5' stay as they are; '+SUM(1)' and '-2+3' get the '). XLSX cells are always typed as text,
    so the Excel file needs (and gets) no prefix at all."""
    if not isinstance(v, str) or not v: return v
    if HANDLE.fullmatch(v): return " " + v      # '@gachamber' (social handle): a leading space keeps it text without a stray '
    if v[0] in "=@\t\r": return "'" + v
    if v[0] in "+-":
        if PLAINNUM.fullmatch(v): return v          # '-5', '+3.5': a number either way, nothing changes
        if NUMBERISH.fullmatch(v) and re.search(r"[\s()]", v) and not re.search(r"\d\s+[+-]\s*\d|\s[+-]\s", v): return v   # '+1 330-555-0181' stays; '+1 -1' is arithmetic
        return "'" + v                              # '+1-330-555-0182', '+1-1', '+SUM(1)': Excel WOULD calculate it
    return v

def _when(t): return datetime.datetime.fromtimestamp(t).strftime("%Y-%m-%d %H:%M") if t else ""
def _slug(s, n=40): return re.sub(r"[^A-Za-z0-9]+", "-", s or "").strip("-")[:n].strip("-")
def _bname(label):
    """Batch name for file names. Round 8: names are stored as typed; labels from before round 8 had ' · <date/time> · <initials>'
    appended by the page, and only that old suffix is dropped."""
    return re.sub(r"(?: · [^·]*\d{1,2}:\d{2}\s*[AP]M)(?: · [A-Z]{1,4})?$", "", label or "")
def _fname(kind, status, ext, batch_label=""):
    tag = {"approved_new": "new", "approved": "approved", "all": "all", "pending": "pending"}.get(status, _slug(status))
    if batch_label: tag += "-" + _slug(_bname(batch_label), 28)
    kind = {"contacts-for-MC-Trade": "MCTrade"}.get(kind, kind)
    # v0.4.6 #7: seconds in the name (a second export in the same minute no longer gets a confusing '-2'), and the MC Trade
    # import file says so: IMPORT-ME_… (its companions are NOT-FOR-IMPORT_…, see comp_name)
    fn = ("IMPORT-ME_" if kind == "MCTrade" else "") + f"GAC-{kind}-{tag}-{datetime.datetime.now().strftime('%Y-%m-%d_%H%M%S')}.{ext}"
    # round 15 (manager r14 M-1): two exports never share a file name ("-2", "-3" ... when the name was used before)
    try:
        used = {r[0] for r in U.db().execute("SELECT filename FROM exports WHERE filename LIKE ?", (fn.rsplit(".", 1)[0] + "%",))}
    except Exception: used = set()
    n = 2; base, dot = fn.rsplit(".", 1)
    while fn in used: fn = f"{base}-{n}.{dot}"; n += 1
    return fn
def plural(n, one, many=None): return f"{n} {one if n == 1 else (many or one + 's')}"
def ids3(ids, k=3):
    """Round 20: '#4, #9, #12 and 7 more' (never a wall of ids)."""
    ids = list(ids); h = ", ".join("#" + str(i) for i in ids[:k])
    return h + (f" and {len(ids) - k} more" if len(ids) > k else "")

def same_names(sel, contacts_from=None):
    """Round 15 (mctrade r14 M1): rows of one export that share a first + last name with another row of the export or with
    an MC Trade contact (real export only). MC Trade's import may treat same first+last as one duplicate: review first."""
    by = {}
    for c in sel:
        n = U._nm(c["fields"])
        if n and " " in n: by.setdefault(n, []).append(c["id"])
    out = [{"name": n.title(), "cards": ids, "with": "export"} for n, ids in by.items() if len(ids) > 1]
    if contacts_from is not None:   # round 20: the MC Trade match looks at every exported card, the 'twice' check at file rows
        by = {}
        for c in contacts_from:
            n = U._nm(c["fields"])
            if n and " " in n: by.setdefault(n, []).append(c["id"])
    rows = crm_source.real()
    if rows:
        have = {}
        for r in rows:
            n = re.sub(r"\s+", " ", (r.get("name") or f"{r.get('first_name','')} {r.get('last_name','')}").strip().lower())
            if n: have.setdefault(n, []).append(str(r.get("profile_id") or r.get("id") or ""))
        for n, ids in by.items():
            if n in have: out.append({"name": n.title(), "cards": ids, "with": "contacts", "profiles": have[n][:5]})
    return out

def _check_ids(sel, metas, status):
    res = build_export("xlsx", sel, metas, status, "summary", "", preview=True)
    sm = res[4] if len(res) > 4 else {}
    dz = {str(k): v for k, v in (sm.get("decisions") or {}).items()}
    return set(sm.get("check_ids") or []) | {x[0] for x in (sm.get("_view") or {}).get("left") or [] if dz.get(str(x[0])) == "review"}

@U.router.get("/api/export/check")
def export_check(request: Request, status: str = "approved_new", batch: str = ""):
    if _filter_refused(status, request.state.user):
        return JSONResponse({"error": "Only an admin can export pending, rejected or all cards.", "admin_only": True}, 403)
    if _batch_refused(batch, request.state.user): return JSONResponse({"error": BATCH_MSG, "admin_only": True}, 403)
    sel, metas, probs = _export_set(status, batch)
    flagged = [c["id"] for c in sel if (metas.get(c["id"]) or {}).get("override") or U.validate(c["fields"])]
    rows_, held_, sc_ = None, [], None
    if sel and len(sel) <= 500:   # v0.4.2 #4: ONE source: the page line uses the sheet's own "to check" list
        try: sc_ = counts_for(sel, metas, status, _inc(request), waiting=waiting_ids(batch) if batch else None); flagged = sc_["check_ids"]; rows_ = sc_.get("rows"); held_ = (sc_.get("changed_ids") or []) + (sc_.get("repeat_ids") or [])
        except Exception: sc_ = None
    ok_batches = []
    if probs and not batch:   # all batches is blocked: say which single batches can be exported right now
        for b in sorted({(metas.get(c["id"]) or {}).get("batch") or "__none__" for c in sel}):
            s2, m2, p2 = _export_set(status, b)
            if s2 and not p2: ok_batches.append({"batch": b, "label": (m2.get(s2[0]["id"]) or {}).get("batch_label") or b, "count": len(s2)})
    cs_, metas_ = U._active()
    pend = [{"id": c["id"], "name": f"{c['fields'].get('first_name','')} {c['fields'].get('last_name','')}".strip() or "(no name)"}
            for c in cs_ if c["status"] == "pending" and (not batch or ((metas_.get(c["id"]) or {}).get("batch") or "__none__") == batch)]
    # v0.4.9 #1: a batch whose approved cards all went out already says so (not "Approve some cards first")
    sent_before = 0
    if not sel and status == "approved_new":
        sid_ = sent_ids()
        sent_before = sum(1 for c in cs_ if c["status"] == "approved" and c["id"] in sid_ and (not batch or ((metas_.get(c["id"]) or {}).get("batch") or "__none__") == batch))
    why_ = (f"{plural(sent_before, 'card')} already exported. Nothing new to import.") if sent_before else (why_line(sc_, len(sel)) if sc_ or not sel else "")
    return {"count": len(sel), "rows": rows_, "held": len(held_), "sent_before": sent_before, "ready": (sc_ or {}).get("ready"), "why": why_, "duplicates": probs, "flagged": flagged, "same_names": same_names(sel), "ok_batches": ok_batches, "pending": pend[:200], "pending_count": len(pend),
            "cards": [{"id": c["id"], "name": f"{c['fields'].get('first_name','')} {c['fields'].get('last_name','')}".strip()} for c in sel[:300]],
            "admin_only": _filter_refused(status, request.state.user), "versions": export_versions(sel)}

def strict_profile_id(c, existing=None):
    """MC Trade Profile ID for an 'update existing' card, ONLY when certain (round 6, R3-8): the card's email matches the
    profile exactly, or the exact same name + the same phone with a high score. Two different certain candidates, or
    anything weaker (e.g. a similar name + a shared office line) -> None, and the card goes to review."""
    sure = []
    existing = existing if existing is not None else crm_source.real()
    if existing is None: return None      # round 7: never a Profile ID from SAMPLE data
    for m_ in U._crm_index(existing).match(c["fields"]):
        exact_name = mctrade_sync._nname(m_["name"]) == mctrade_sync._nname(f"{c['fields'].get('first_name','')} {c['fields'].get('last_name','')}")
        if "email exact" in m_["why"] or (exact_name and "phone exact" in m_["why"] and m_["score"] >= 0.5): sure.append(m_["profile_id"])
    sure = list(dict.fromkeys(sure))
    return sure[0] if len(sure) == 1 else None
def exact_email_profile(c, existing=None):
    """Profile ID of an existing MC Trade profile with exactly this card's email (create-new cards: would duplicate it)."""
    existing = existing if existing is not None else crm_source.real()
    if existing is None: return None
    hits = list(dict.fromkeys(m_["profile_id"] for m_ in U._crm_index(existing).match(c["fields"]) if "email exact" in m_["why"]))
    return hits[0] if hits else None

def crm_look(c, existing=None):
    """Round 8: everything classify() needs about one card vs the REAL MC Trade contacts (see mctrade_export.classify)."""
    existing = existing if existing is not None else crm_source.real()
    if existing is None: return {}
    f = c["fields"]; nm = mctrade_sync._nname(f"{f.get('first_name','')} {f.get('last_name','')}")
    byid = U._crm_index(existing).byid
    cem = {e.strip().lower() for e in f.get("emails") or [] if (e or "").strip()}
    email, strong, name = [], [], []
    for m_ in U._crm_index(existing).match(f):
        pid = str(m_["profile_id"]); exact = bool(nm) and mctrade_sync._nname(m_["name"]) == nm
        if "email exact" in m_["why"]: email.append(pid); strong.append(pid)
        elif exact and ("phone exact" in m_["why"] or any(w.startswith("org~") for w in m_["why"])): strong.append(pid)
        elif exact: name.append(pid)
    email, strong, name = (list(dict.fromkeys(x)) for x in (email, strong, name))
    sure = None if len(email) > 1 else strict_profile_id(c, existing)
    pe = ((byid.get(str(sure)) or {}).get("email") or "").strip().lower() if sure else ""
    return {"email": email, "strong": strong, "name": name, "sure": sure,
            # round 11 (mctrade-r7 F1): the UPDATES row writes the card's FIRST email; any difference from the profile's email
            # holds the card back (a match on the card's second email would otherwise overwrite the profile email)
            "email_differs": bool(sure and pe and (f.get("emails") or [""])[0].strip().lower() not in ("", pe))}

def _note(c):
    """Note as exported: a stale 'Duplicate of #N (details merged)' on a card that is no longer rejected is dropped."""
    n = c.get("note") or ""
    return U.STALE_DUP.sub("", n).strip(" |") if c.get("status") != "rejected" else n
def last_export_file(cid):
    """The file a card was last exported in ('' if none)."""
    r = U.db().execute("SELECT detail FROM card_history WHERE card_id=? AND action='exported' ORDER BY id DESC LIMIT 1", (cid,)).fetchone()
    try: return (json.loads(r[0] or "{}") or {}).get("file") or "" if r else ""
    except Exception: return ""
# v0.4.8 #1-3: one plain sentence per reason a card is not in the import file
CHANGED_WHY = "Edited after export, so it is not in this file unless an admin includes it."   # v0.4.9 #6: one wording
REPEAT_WHY = "Already exported, so it is not in this file."
CSV_ONLY = "Sent before only in a contacts list (CSV), not to MC Trade."
def _norm(s): return re.sub(r"[^a-z0-9@.]+", " ", str(s or "").lower()).strip()
def _nm(s): return re.sub(r"[^a-z0-9]+", "", mctrade_sync.fold(str(s or "")).lower())
def person_keys(first, last, company, emails):
    """v0.4.8 #2: one person = the same email, or the same first + last name at the same company.
    v0.4.9 #3: case, punctuation, a '+tag' in the email and company suffixes (Co / Co. / Company / Inc / LLC ...) are ignored."""
    out = {"e:" + mctrade_sync.norm_email(str(e)) for e in emails or [] if "@" in str(e or "")}
    nm = _nm(f"{first or ''}{last or ''}"); co = mctrade_sync.norm_company(str(company or ""))
    if nm and co: out.add("n:" + nm + "|" + co)
    return out
def _canon(k):
    """A key stored by an older version, in today's form (so 'Test Co.' in an old file matches 'Test Co' today)."""
    if k.startswith("e:"): return "e:" + mctrade_sync.norm_email(k[2:])
    if k.startswith("n:") and "|" in k:
        nm, co = k[2:].split("|", 1); return "n:" + _nm(nm) + "|" + mctrade_sync.norm_company(co)
    return k
def card_keys(c):
    f = c["fields"]; return person_keys(f.get("first_name"), f.get("last_name"), f.get("company"), f.get("emails") or [])
_PEOPLE_CACHE = {}
_IDX_CACHE = {"sig": None, "val": None}
def _file_people(path):
    """Person keys of the rows of an earlier import file (older exports didn't store them)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True); ws = wb.active; it = ws.iter_rows(values_only=True)
    hdr = [str(h or "") for h in next(it, [])]; ix = {h: i for i, h in enumerate(hdr)}
    g = lambda r, k: (r[ix[k]] if k in ix and ix[k] < len(r) else "") or ""
    out = set()
    for r in it:
        out |= person_keys(g(r, "IndFirstName"), g(r, "IndLastName"), g(r, "OrgName"), [g(r, "IndEmail"), g(r, "IndEmail2")])
    wb.close(); return out
def import_index():
    """v0.4.9 #2: THE 'Sent to MC Trade' record. Only MC Trade import files (IMPORT-ME, xlsx) count; a contacts list (CSV)
    never does. -> (people {key: export}, cards {card id: export}); export = {id, file, ts, user, batch, ids, keys}.
    Exports made before v0.4.8 are read from their stored file (or, when the copy is gone, from the cards they listed)."""
    sig = tuple(U.db().execute("SELECT COUNT(*), MAX(id) FROM exports WHERE kind='xlsx'").fetchone())
    if _IDX_CACHE["sig"] == sig and _IDX_CACHE["val"] is not None: return _IDX_CACHE["val"]
    people, cards = {}, {}
    rows = U.db().execute("SELECT id, ts, username, batch_label, filename, card_ids, stored, people FROM exports WHERE kind='xlsx' ORDER BY id").fetchall()
    cs = None
    for r in rows:
        ex = {"id": r["id"], "file": r["filename"], "ts": r["ts"], "user": r["username"], "batch": r["batch_label"] or ""}
        try: ids = [int(i) for i in json.loads(r["card_ids"] or "[]")]
        except Exception: ids = []
        keys = None
        if r["people"]:
            try: keys = {_canon(k) for k in json.loads(r["people"])}
            except Exception: keys = None
        if keys is None:
            if r["id"] in _PEOPLE_CACHE: keys = _PEOPLE_CACHE[r["id"]]
            else:
                p = os.path.join(U.EXP, r["stored"]) if r["stored"] else ""
                try: keys = _file_people(p) if p and os.path.isfile(p) else None
                except Exception: keys = None
                if keys is None:
                    if cs is None: cs = {c["id"]: c for c in store.all_()}
                    keys = set().union(*[card_keys(cs[i]) for i in ids if i in cs]) if cs else set()
                _PEOPLE_CACHE[r["id"]] = keys
        if not r["people"] and keys:   # before v0.4.8 card_ids also listed cards left out of the file: keep only the people in it
            if cs is None: cs = {c["id"]: c for c in store.all_()}
            ids = [i for i in ids if i not in cs or not card_keys(cs[i]) or card_keys(cs[i]) & keys]
        ex["ids"] = set(ids); ex["keys"] = keys
        for k in keys: people.setdefault(k, ex)
        for i in ids: cards[i] = ex   # the LATEST import file a card was in
    _IDX_CACHE.update(sig=sig, val=(people, cards))
    return people, cards
def sent_ids():
    """Card ids that were in an MC Trade import file (the card itself, not just the same person)."""
    return set(import_index()[1])
def sent_record(c, metas, idx=None):
    """v0.4.9 #2/#5: the earlier MC Trade import file this card or this PERSON was in, or None.
    -> (export, own): own = this very card was in it. A look-alike the reviewer marked 'Different person' doesn't count."""
    people, cards = idx if idx is not None else import_index()
    if c["id"] in cards: return cards[c["id"]], True
    kept = U._kept(metas.get(c["id"]) or {})
    for k in sorted(card_keys(c)):
        ex = people.get(k)
        if not ex: continue
        same = {i for i in ex["ids"] if i != c["id"]}
        if same and same <= (kept | {i for i in same if c["id"] in U._kept(metas.get(i) or {})}): continue
        return ex, False
    return None, False
def held_cards(sel, metas, idx=None):
    """v0.4.8 #1/#2: cards that would be a second 'create new' in MC Trade, so they are not in the import file:
    'changed' = THIS card was in an import file and was edited since; 'repeat' = this card or this person is already in an
    earlier import file. -> (changed [(card, export)], repeat [(card, export)], csv_only {card id}).
    v0.4.9 #2: only import files count; a card sent only in a contacts list (CSV) is csv_only, never 'edited after export'."""
    idx = idx if idx is not None else import_index()
    chg, rep, csv_only = [], [], set()
    for c in sel:
        m = metas.get(c["id"]) or {}
        if m.get("crm_action") == "update_existing": continue   # an update of an existing profile is not a new row
        ex, own = sent_record(c, metas, idx)
        if ex: (chg if own and m.get("reexport") else rep).append((c, ex))
        elif m.get("exported"): csv_only.add(c["id"])
    return chg, rep, csv_only
def when_where(ex, user=None):
    """v0.4.9 #4: 'in an earlier export (Oct 28)'; the file name only for the viewer's OWN export."""
    t = datetime.datetime.fromtimestamp(ex["ts"]) if ex and ex.get("ts") else None
    d = f"{t:%b} {t.day}" if t else ""   # no "%-d": it is glibc-only and raises on Windows
    own = bool(user and ex and (ex.get("user") or "") == (user.get("username") if isinstance(user, dict) else user))
    return (f"in {ex['file']} ({d})" if own and ex.get("file") else f"in an earlier export ({d})" if d else "in an earlier export")
def changed_after_export(sel, metas):
    """v0.4.7 #2 (kept for callers): cards exported in an import file and edited since."""
    return [c for c, _ in held_cards(sel, metas)[0]]

def build_export(kind, sel, metas, status, who, batch_label="", preview=False, include_changed=False, waiting=None, vis=None, viewer=None):
    def flag(c):
        m = metas.get(c["id"]) or {}; ov = U._ov(m); errs = U.validate(c["fields"]); out = []
        if ov: out.append(f"APPROVED ANYWAY by {ov.get('by')}: {ov.get('reason')}")
        if errs: out.append("CHECK: " + "; ".join(errs.values()))
        out += U.confirmed_lines(c["fields"])   # round 22: hard-check overrides are visible to the importer
        xb = U.db().execute("SELECT detail FROM card_history WHERE card_id=? AND action='cross_batch_confirmed' ORDER BY id DESC LIMIT 1", (c["id"],)).fetchone()
        if xb:
            try: oth = (json.loads(xb[0] or "{}") or {}).get("cards") or []
            except Exception: oth = []
            # v0.4.2 #3 (data-manager): look-alikes later answered 'Different' are not a double-import risk
            kp_ = set(U._kept(m)); oth = [o for o in oth if o not in kp_ and c["id"] not in set(U._kept(metas.get(o) or {}))] if oth else oth
            lst_ = _oth_list(c["id"], oth) if oth else ""
            if xb_ok(c, oth): out.append("Also in another batch: approved over a look-alike" + (f" ({lst_})" if oth else "") + ": check for a double import")
        return " | ".join(out)
    def _oth_list(cid, oth):   # v0.4.8 #4: a reviewer sees only same-batch numbers, the rest as a count
        if vis is None: return ", ".join("#" + str(o) for o in oth)
        bx_ = (metas.get(cid) or {}).get("batch"); sh_ = [o for o in oth if bx_ and o in vis and (metas.get(o) or U.meta(o) or {}).get("batch") == bx_]
        nh_ = len(oth) - len(sh_)
        return ", ".join(["#" + str(o) for o in sh_] + ([f"{nh_} card{'s' if nh_ != 1 else ''} {U._ob(nh_)}"] if nh_ else []))
    def xb_ok(c, oth):
        m = metas.get(c["id"]) or {}
        if oth: return True
        return not (U._kept(m) or m.get("dup_resolution") == "keep_both")
    def aud(c):
        m = metas.get(c["id"]) or {}
        return [c["id"], c.get("reviewer") or "", _when(c.get("decided")), m.get("batch_label") or "",
                {"update_existing": "update existing", "create_new": "create new"}.get(m.get("crm_action"), "create new (no MC Trade match)" if c["status"] == "approved" else ""), flag(c)]   # round 13 (newcomer r11 #6): never blank for an exported new profile
    def more(c):
        f = c["fields"]
        return [f.get("credentials") or "", f.get("department") or "", f.get("other_address") or "",
                "; ".join(f.get("emails") or []), "; ".join(f"{p.get('type')}: {p.get('number')}" for p in f.get("phones") or []), _note(c)]
    if kind == "csv":
        b = io.StringIO(); w = csv.writer(b)
        w.writerow(export.COLS + ["credentials", "department", "other_address", "all_emails", "all_phones", "note",
                                  "reviewed_by", "decided_at", "batch", "mc_trade_action", "needs_attention"])
        for c in sel:
            r = export._flat(c); r["social"] = r["social"]; a = aud(c); w.writerow([safe_cell(x) for x in [r[k] for k in export.COLS] + more(c) + a[1:]])
        return ("\ufeff" + b.getvalue()).encode(), "text/csv", _fname("contacts", status, "csv", batch_label)   # BOM: Excel opens UTF-8 correctly
    if kind == "vcf":
        return export.to_vcard(sel).encode(), "text/vcard", _fname("contacts", status, "vcf", batch_label)
    if kind == "xlsx":   # MC Trade import file: template columns only, one sheet, one row per person (see mctrade_export.py)
        from openpyxl import Workbook
        from . import mctrade_export as mx
        aff = auth.get_setting("mctrade_affcode") or ""; style = auth.get_setting("phone_format") or "gac"
        def _kept_pair(a, b_): return b_ in U._kept(metas.get(a) or {}) or a in U._kept(metas.get(b_) or {})
        _real = crm_source.real(); _hint_rows, _ = crm_source.rows_for_hints()
        def _hint(c):
            m_ = U._crm_index(_hint_rows).match(c["fields"])
            return f"{m_[0]['name']} ({m_[0]['organization']})" if m_ else None
        # v0.4.8 #1/#2: cards already sent to MC Trade (edited since, or unchanged / the same person again) are not new rows
        _chgp, _repp, _csv_only = held_cards(sel, metas)
        _held_all = [(c, ex, "changed") for c, ex in _chgp] + [(c, ex, "repeat") for c, ex in _repp]
        _held = [] if include_changed else _held_all
        _chg = [c for c, _x, k in _held if k == "changed"]; _rep = [c for c, _x, k in _held if k == "repeat"]
        _chg_ids = {c["id"] for c in _chg}; _rep_ids = {c["id"] for c in _rep}; _held_ids = _chg_ids | _rep_ids
        # v0.4.9 #4: an earlier export is named by its date; the file name only when it is the viewer's own export
        _chg_files = {c["id"]: when_where(_x, viewer) for c, _x, _k in _held}
        res = mx.build([c for c in sel if c["id"] not in _held_ids], aff, style, U.hard_errors, kept=_kept_pair, confirmed=lambda f_: list(U.valid_confirmations(f_)), crm_look=lambda c: crm_look(c, _real),
                       crm_real=_real is not None, crm_hint=_hint, updates_on=auth.get_setting("mctrade_updates_enabled") == "1",
                       action=lambda c: (metas.get(c["id"]) or {}).get("crm_action") or "")
        for c, _x, k in _held:   # in neither import file; listed in the README, the summary and the CSV
            w_ = CHANGED_WHY if k == "changed" else REPEAT_WHY
            res["left"].append((c, w_)); res["decisions"][c["id"]] = (mx.LEFT, w_)
        hdr, source, rows, left, upd = res["headers"], res["source"], res["rows"], res["left"], res["updates"]
        def _book(rws):
            wb = Workbook(); ws = wb.active; ws.title = "Sheet1"; ws.append(hdr)
            for c, r_ in rws: ws.append(mx.cells(hdr, r_))
            for rw in ws.iter_rows():
                for cell in rw:
                    if isinstance(cell.value, str): cell.data_type = "s"   # plain text, never a formula; no ' prefix needed
            b_ = io.BytesIO(); wb.save(b_); return b_.getvalue()
        body_x = _book(rows)
        fn = _fname("contacts-for-MC-Trade", status, "xlsx", batch_label)
        rfn = preview_name(fn) if preview else fn   # round 10 (coordinator r3 3c): the READ-ME names the files as they are saved
        # companion files (not for import): every non-template column + rows left out, and the READ ME notes
        e = io.StringIO(); w = csv.writer(e)
        w.writerow(["ScanCardID", "InImportFile", "Why left out", "FirstName", "LastName", "Email2", "Website", "Credentials", "Department", "OtherAddress",
                    "AllEmails", "AllPhones", "Social", "Note", "ReviewedBy", "DecidedAt", "Batch", "MCTradeAction", "Check"])
        why = {c["id"]: r for c, r in left}; updids = {c["id"] for c, _ in upd}
        cinfo = crm_source.info()
        _smp = []   # round 9 (newcomer r9 N1): cards in the new file that match a SAMPLE contact by email (SAMPLE data only)
        if not cinfo["real"]:
            _sr = crm_source.sample_rows(); _in = {x["id"] for x, _ in rows}
            _smp = [c_["id"] for c_ in sel if c_["id"] in _in and any("email exact" in m_["why"] for m_ in mctrade_sync.match(c_["fields"], _sr))]
        # v0.4.3 #4 (mctrade-admin L3): a possible MC Trade match is a note in the CSV Check column too
        _mcn = {i: "May already be in MC Trade (practice contacts only)" for i in _smp} | {i: "May already be in MC Trade" for i in res["possible"] if i not in _smp}   # v0.4.8 #8: capital M
        def aud(c, _a=aud):
            a_ = _a(c)
            if c["id"] in _mcn and _mcn[c["id"]] not in a_[-1]: a_[-1] = " | ".join(x for x in (a_[-1], _mcn[c["id"]][0].upper() + _mcn[c["id"]][1:]) if x)
            return a_
        nflag = sum(1 for c in sel if aud(c)[-1])
        _imp_ids = {c["id"] for c, _ in rows} | {c["id"] for c, _ in upd}
        _flagged = {c["id"] for c in sel if aud(c)[-1]}
        xb_ids = {}
        for c in sel:   # cards approved over a cross-batch look-alike warning (coordinator r1 #6): risk of a double import
            for h in U.history(c["id"]):
                if h["action"] == "cross_batch_confirmed":
                    d_ = h.get("detail") if isinstance(h.get("detail"), dict) else {}
                    xb_ids[c["id"]] = d_.get("cards") or []
            # v0.4.1 #5 (datamanager #4): 'Different' reaches the export: look-alikes marked different are not a double-import risk
            kp_ = set(U._kept(metas.get(c["id"]) or {}))
            if c["id"] in xb_ids:
                xb_ids[c["id"]] = [o for o in xb_ids[c["id"]] if o not in kp_ and c["id"] not in set(U._kept(metas.get(o) or {}))]
                if not xb_ids[c["id"]] or (metas.get(c["id"]) or {}).get("dup_resolution") == "keep_both" and not kp_: xb_ids.pop(c["id"])
        tpl_real = source == "GAC template file"
        cnt_ = mx.counts(res["decisions"])   # round 22 (mctrade-admin r21 H1): ONE source for every count (README, sign-off, dry run)
        conf_ = [(c["id"], ln) for c in sel for ln in U.confirmed_lines(c["fields"])]
        ovr_ = [c["id"] for c in sel if U._ov(metas.get(c["id"]) or {})]
        # v0.3.2 (newcomer v0.3.1): short and plain; empty/zero lines left out; ONE warning line per card
        _warn = {}
        for i in _smp: _warn.setdefault(i, []).append("may already be in MC Trade (practice contacts only)")
        for i in res["possible"]:
            if i not in _smp: _warn.setdefault(i, []).append("may already be in MC Trade")
        for i in xb_ids: _warn.setdefault(i, []).append("may match someone approved in another batch")
        if include_changed:
            for c_, _x, _k in _held_all: _warn.setdefault(c_["id"], []).append("already exported: may be added twice")
        _in_rows = {c_["id"] for c_, _ in rows}
        for i in sorted(_csv_only & _in_rows): _warn.setdefault(i, []).append(CSV_ONLY[0].lower() + CSV_ONLY[1:].rstrip("."))
        _cap = lambda t_: t_[:1].upper() + t_[1:]
        _wl = [f"Check #{i}: {_cap(' and '.join(v))}." for i, v in sorted(_warn.items())]
        # v0.4.9 #2: the CSV "Check" column carries every note the READ-ME points to ('see the CSV'), the 'Already exported'
        # and contacts-list-only notes included
        _hw = {c_["id"]: (CHANGED_WHY if k_ == "changed" else REPEAT_WHY).replace(", so it is not in this file", f" ({_chg_files.get(c_['id']) or 'in an earlier export'}), so it is not in this file")
               for c_, _x, k_ in _held}
        for c in sel:
            f_ = c["fields"]; a = list(aud(c)); em = f_.get("emails") or []
            notes_ = [x for x in str(a[-1] or "").split(" | ") if x]
            for t_ in [_cap(x) for x in _warn.get(c["id"], [])] + ([_hw[c["id"]].rstrip(".")] if c["id"] in _hw else []) + ([CSV_ONLY.rstrip(".")] if c["id"] in _csv_only and c["id"] not in _in_rows else []):
                if t_.lower() not in " | ".join(notes_).lower(): notes_.append(t_)
            a[-1] = " | ".join(notes_)
            inimp = "updates file" if c["id"] in updids else "no" if c["id"] in why else "yes"
            w.writerow([safe_cell(x) for x in [c["id"], inimp, why.get(c["id"], ""), mx.clean(f_.get("first_name")), mx.clean(f_.get("last_name")),
                        em[1] if len(em) > 1 else "", mx.clean(f_.get("website"))] + [mx.clean(x) for x in more(c)[:5]] + [" ".join(f_.get("social") or []), mx.clean(_note(c))] + a[1:]])
        _named = set(_warn) | {i for i, _ in conf_[:20]} | set(ovr_) | {i for i, _w, _y in website_warnings(sel, metas, rows + upd, res.get("merged"))} \
            | {k for k, v in res["decisions"].items() if v[0] == mx.REVIEW}
        _more = _flagged - _named
        # v0.4.4 #3 (data-manager M1, reviewer): ONE 'To check' set: the Check-line cards + the cards a person must decide (NEEDS REVIEW).
        # Screen, Save dialog, READ-ME Totals, dry run and Preview all use it.
        _chk0 = set(_warn) | {i for i, _ in conf_} | set(ovr_) | _flagged | {i for i, _w, _y in website_warnings(sel, metas, rows + upd, res.get("merged"))}
        _tc = _chk0 | {int(k) for k, v in res["decisions"].items() if v[0] == mx.REVIEW}
        # v0.4.5 #8 (data-manager: Why said only "see the READ-ME"): the real reasons, so the summary works on its own
        _why = {i: list(v) for i, v in _warn.items()}
        _cl = {}   # v0.4.7 #9 (data-manager): short why, "Email, website confirmed" (the full lines stay in the READ-ME)
        for i, ln in conf_: _cl.setdefault(i, []).append(str(ln).split(" confirmed by ")[0].strip())
        for i, labs in _cl.items():
            labs = list(dict.fromkeys(labs)); t_ = ", ".join([labs[0]] + [x[:1].lower() + x[1:] for x in labs[1:]]) + " confirmed"
            _why.setdefault(i, []).append(t_)
        for i in ovr_: _why.setdefault(i, []).append("approved with a reason despite a warning")
        for i, w_, y_ in website_warnings(sel, metas, rows + upd, res.get("merged")): _why.setdefault(i, []).append(f"website {w_}: {y_}")
        for k, v in res["decisions"].items():
            if v[0] == mx.REVIEW: _why.setdefault(int(k), []).append("needs a decision: " + str(v[1] or "possible match in MC Trade").rstrip("."))
        for c in sel:
            if c["id"] in _flagged and c["id"] not in _why: _why[c["id"]] = [str(aud(c)[-1]).rstrip(".")]
        _has_web = any(mx.canonical(h) in ("Website", "IndWebsite") for h in hdr if h)
        # v0.4.6 #7: ONE 'Ready to import' line (screen, READ-ME, summary sheet, CSV)
        # v0.4.7 #3: rows = clean + to_check (in the file); cards to check that are NOT in the file are counted apart
        _merged_in = {x for t_, ids_ in (res.get("merged") or {}).items() if t_ in _imp_ids for x in ids_}
        _in_file = _imp_ids | _merged_in   # v0.4.8 #1: only these are marked exported (and listed in the export's card_ids)
        _fix = {c["id"] for c, w_ in left if str(w_).startswith("FIX FIRST")}
        _waiting = [i for i in (waiting or []) if i not in _in_file]
        _ready = {"rows": len(rows) + len(upd), "clean": len(_imp_ids - _tc), "to_check": len(_imp_ids & _tc),
                  "outside": len((_tc | {int(k) for k, v in res["decisions"].items() if v[0] == mx.REVIEW}) - _in_file - _held_ids),
                  "changed": len(_chg), "repeat": len(_rep), "fix": len(_fix - _held_ids), "waiting": len(_waiting),
                  "files": (4 + bool(upd)) if (rows or upd) else 0}
        # v0.4.7 #9 (data-manager L5): no summary row at the end of the CSV (it read as a blank row); the Ready line is in the READ-ME and summary
        readme = "\n".join([
            *(["*** PREVIEW: import none of these files. Nothing was marked as exported; use Export for the real file. ***"] if preview else []),
            "GAC Card Scanner: MC Trade import file" if not preview else "GAC Card Scanner: MC Trade preview (not for import)",
            ready_line(_ready) if not preview else "",
            readme_totals(cnt_, len(_imp_ids & _tc), 0, notin=[] if not preview else notin_parts(_ready)),
            *([f"{ids3(sorted(_chg_ids))}: " + (CHANGED_WHY if len(_chg) == 1 else CHANGED_WHY.replace("so it is", "so they are").replace("includes it", "includes them"))
               + (f" Sent {sorted(set(_chg_files[i] for i in _chg_ids))[0]}." if len(set(_chg_files[i] for i in _chg_ids)) == 1 else "")] if _chg else []),   # v0.4.9 #6: one wording
            *([f"{ids3(sorted(_rep_ids))}: Already exported, so not in this file"
               + (f", {sorted(set(_chg_files[i] for i in _rep_ids))[0]}" if len(set(_chg_files[i] for i in _rep_ids)) == 1 else "") + "."] if _rep else []),
            *([f"{ids3(_waiting)}: Waiting for an admin, so not in this file."] if _waiting else []),
            ("" if preview else ("Before importing: read the “Check” lines. " if _tc else "") + "An import can't be undone."
             + (" Import new contacts and updates separately." if upd else "")),
            f"Made by GAC Card Scanner {_ver()} on {_when(time.time())} by {who}" + (f" · batch {batch_label}" if batch_label else "") + f" · {plural(len(sel), 'card')}.",
            # v0.4.6 #7 (READ-ME said "DO NOT IMPORT" and "Import this file only"): a preview names no file to import
            (f"Import only: {rfn}" + (f" (and {comp_name(rfn, '-UPDATES.xlsx')} for updates)" if upd else "") + ". The NOT-FOR-IMPORT files are for reading."
             if not preview else "These are preview copies: import none of them."),
            *[f"Check: #{i} {ln}" for i, ln in conf_[:20]],
            (f"… and {len(conf_) - 20} more confirmed values in the CSV." if len(conf_) > 20 else ""),
            (f"Check: {plural(len(ovr_), 'card')} approved anyway (see the CSV)." if ovr_ else ""),
            ("Compared with MC Trade contacts: " + cinfo["source"] + "." if cinfo["real"] else
             "No list of MC Trade contacts is loaded, so the scanner could not check who is already in MC Trade."),
            (f"Updates: {comp_name(rfn, '-UPDATES.xlsx')} ({plural(len(upd), 'row')}). MemberSince is left blank so it can't overwrite a member's real date."
             if upd else ""),
            (f"Template columns the scanner leaves blank: {', '.join(res['unknown_cols'])}." if res["unknown_cols"] else
             "" if tpl_real else "The column names have not been checked against MC Trade's own import template: test 2 rows first."),
            (f"Not in this template (CSV only): {', '.join(res['unfilled'])}." if res["unfilled"] else ""),
            (f"Every row gets the group code '{aff}'. It must already exist in MC Trade." if aff else ""),
            ("" if _has_web else "Websites: this template has no website column, so websites are only in the CSV."),
            (f"Combined with another row: {ids3(sorted(_merged_in))}." if _merged_in else ""),
            (f"To fix first, not in this file: {ids3(sorted(_fix - _held_ids))}. The CSV says why." if _fix - _held_ids else ""),
            # round 20 (mctrade-admin r19 F1): 'in this export N times' only when 2+ ROWS of the import files share the name
            *[f"Check, same name: {g['name']} is on {plural(len(g['cards']), 'card')} {ids3(g['cards'])}"
              + (f" ({plural(len(g['cards']), 'row')} in this export)" if g["with"] == "export" else f" and already in MC Trade (profile {', '.join(g.get('profiles') or [])})")
              + ". MC Trade may merge people with the same name: check they are different, or merge them." for g in same_names([c for c, _ in rows] + [c for c, _ in upd], contacts_from=sel)],
            (f"Check, to review: {cnt_[mx.REVIEW]} {U._s(cnt_[mx.REVIEW], 'card')} ({ids3([k for k, v in res['decisions'].items() if v[0] == mx.REVIEW])}) "
             "are in neither file: a person must decide (the CSV says why)." if cnt_[mx.REVIEW] else ""),
            *_wl[:10],
            (f"… and {len(_wl) - 10} more cards to check (see the CSV)." if len(_wl) > 10 else ""),
            *(["Check, marked different but identical: " + x for x in res["conflicts"]] if len(res["conflicts"]) <= 3 else
              [f"Check, marked different but identical: {len(res['conflicts'])} cards, each exported once (the CSV says which)."]),
            *[f"Check website: #{i} has {w} ({why}). Compare it with the photo." for i, w, why in website_warnings(sel, metas, rows + upd, res.get("merged"))],
            # v0.4.2 #4 (coordinator L): plain wording; the same cards are in the sheet's "To check" count
            # v0.4.3 #8 (data-manager/newcomer): "More notes" counts only cards NOT already named on a Check line above
            (f"More notes: {plural(len(_more), 'other card')} {'has' if len(_more) == 1 else 'have'} a note in the CSV “Check” column. Read them before importing." if _more else ""),
            "Second emails and phones, notes and reviewer are in the CSV, not in the import file."])
        readme = "\n".join(x for x in readme.split("\n") if x.strip())
        _seen = set(); readme = "\n".join(x for x in readme.split("\n") if not x.strip() or not (x in _seen or _seen.add(x)))   # round 24: each warning once
        comp = {"extra.csv": ("\ufeff" + e.getvalue()).encode(), "readme.txt": readme.encode()}
        if upd: comp["updates.xlsx"] = _book(upd)
        # v0.4.4 #8 (data-manager): the summary sheet is back, as its own workbook (the import file stays one template sheet)
        _sw = Workbook(); _ss = _sw.active; _ss.title = "Summary"
        for rw_ in (["GAC Card Scanner export summary (not for import)"], [ready_line(_ready)], ["Batch", batch_label or "All batches"], ["Cards", len(sel)],
                    ["New contacts", cnt_[mx.NEW]], ["Updates", cnt_[mx.UPDATE]], ["Rows to check", len(_imp_ids & _tc)], ["Cards to decide", _ready["outside"]],
                    *([["Edited after export (not in this file)", len(_chg)]] if _chg else []), *([["Already exported (not in this file)", len(_rep)]] if _rep else []),
                    *([["Waiting for admin", len(_waiting)]] if _waiting else []), [],
                    ["Card", "Why to check it"], *[[f"#{i}", _cap("; ".join(dict.fromkeys(_why.get(i, []))) or (_mcn.get(i) or "a note in the CSV Check column"))] for i in sorted(_tc)],
                    *([[], ["Card", "Not in this file", "Earlier export"]] + [[f"#{c['id']}", CHANGED_WHY if k == "changed" else REPEAT_WHY, _chg_files.get(c["id"]) or ""] for c, _x, k in _held] if _held else [])):
            _ss.append([safe_cell(x) if isinstance(x, str) else x for x in rw_])
        for rw in _ss.iter_rows():
            for cell in rw:
                if isinstance(cell.value, str): cell.data_type = "s"
        _sb = io.BytesIO(); _sw.save(_sb); comp["summary.xlsx"] = _sb.getvalue()
        # round 9 (mctrade-r6): what each update row would overwrite in MC Trade, old value next to new
        _byid = {str(r_["profile_id"]): r_ for r_ in (crm_source.real() or [])}; _udiff = []
        _inhdr = {mx.canonical(h) for h in hdr if h}   # round 13 (mctrade-r8 R2): only columns that are really in the file (no phantom IndWebsite)
        for c_, r_ in upd:
            old_ = _byid.get(str(r_.get("ProfileID"))) or {}; f_ = c_["fields"]
            pairs = [("Name", f"{old_.get('first_name', '')} {old_.get('last_name', '')}".strip(), f"{f_.get('first_name', '')} {f_.get('last_name', '')}".strip()),
                     ("Organization", old_.get("organization", ""), f_.get("company", "")),
                     ("Email", old_.get("email", ""), (f_.get("emails") or [""])[0]),
                     ("Phone", old_.get("phone", ""), ((f_.get("phones") or [{}])[0] or {}).get("number", ""))]
            _udiff.append({"card": c_["id"], "profile_id": str(r_.get("ProfileID")),
                           "fields": [{"field": k, "old": o or "", "new": n or "", "changes": U._digits(o) != U._digits(n) if k == "Phone" else (o or "").strip().lower() != (n or "").strip().lower()} for k, o, n in pairs]
                           # round 11 (mctrade-r7 F2): EVERY other column the UPDATES row writes is listed too. The contacts export
                           # doesn't hold MC Trade's current value for these, so it is shown as unknown (old=None).
                           + [{"field": h_, "old": None, "new": str(v_), "changes": None}
                              for h_, v_ in r_.items() if v_ not in (None, "") and mx.canonical(h_) in _inhdr and h_ not in ("ProfileID", "MemberSince", "OrgInd", "IndFirstName", "IndLastName", "OrgName", "IndEmail", "IndPhone")]})
        _ready["text"] = ready_line(_ready)
        summ = {"ready": _ready, "mc_possible": {str(k): v for k, v in sorted(_mcn.items())}, "check_ids": sorted(_chk0), "to_check_ids": sorted(_tc),   # v0.4.1 #4: the cards behind the READ-ME "Check" lines
                "new_rows": len(rows), "update_rows": len(upd), "cards": len(_in_file), "left_out": len(left), "sample_email_matches": _smp, "changed_ids": sorted(_chg_ids),
                "repeat_ids": sorted(_rep_ids), "held_files": {str(k): v for k, v in _chg_files.items()}, "csv_only_ids": sorted(_csv_only),
                "in_file_ids": sorted(_in_file), "people": sorted(set().union(*[card_keys(c_) for c_ in sel if c_["id"] in _in_file and c_["id"] not in updids])) if _in_file else [],
                "waiting_ids": _waiting,
                "left_ids": [c["id"] for c, _ in left if c["id"] not in _in_file and c["id"] not in _held_ids],
                "review": res["review"], "possible": res["possible"], "crm_real": cinfo["real"],
                "decisions": {str(k): v[0] for k, v in res["decisions"].items()}, "counts": mx.counts(res["decisions"]),
                # v0.4.2 #3 (reviewer Med): the preview says which new rows may already be in MC Trade, instead of a plain "new profile"
                "_view": {"headers": [h for h in hdr], "rows": [mx.cells(hdr, r_) for _, r_ in rows], "row_notes": [(lambda t_: t_[:1].upper() + t_[1:])("; ".join(_warn.get(c["id"], []))) for c, _ in rows], "row_ids": [c["id"] for c, _ in rows], "updates": [mx.cells(hdr, r_) for _, r_ in upd], "update_diff": _udiff,
                          "left": [[c["id"], why_] for c, why_ in left], "source_text": f"Columns: {source}. MC Trade contacts: {cinfo['source']}.",
                          "warnings": [f"#{i}: approved although look-alikes ({_oth_list(i, os_)}) were approved in another batch" for i, os_ in xb_ids.items()] + [f"#{i}: check the website {w} ({why})" for i, w, why in website_warnings(sel, metas, rows + upd, res.get("merged"))]}}
        return body_x, MT["xlsx"], fn, comp, summ
    raise HTTPException(404)

def _ver():
    try:
        from . import updates; return updates.current_version()
    except Exception: return ""
def notin_parts(r):
    """v0.4.8 #1: why cards are not in the file, one short reason each ('1 edited after export')."""
    return [x for x in (f"{r['changed']} edited after export" if r.get("changed") else "", f"{r['repeat']} already exported" if r.get("repeat") else "",
                        f"{r['outside']} to decide" if r.get("outside") else "", f"{r['fix']} to fix first" if r.get("fix") else "",
                        f"{r['waiting']} waiting for admin" if r.get("waiting") else "") if x]
def ready_line(r):
    """v0.4.6 #7 / v0.4.7 #3 / v0.4.8 #1: 'Ready to import: 3 rows (2 clean, 1 to check). Not in the file: 1 edited after export.
    Test 2 rows first.' clean + to check = rows, always; with no rows: 'Nothing new to import.' plus the reason."""
    t = f"Ready to import: {plural(r['rows'], 'row')} ({r['clean']} clean, {r['to_check']} to check)." if r["rows"] else "Nothing new to import."
    p = notin_parts(r)
    if p: t += " Not in the file: " + ", ".join(p) + "."
    return t + (" Test 2 rows first." if r["rows"] > 2 else "")
def preview_name(fn):
    return "PREVIEW-NOT-FOR-IMPORT_" + re.sub(r"^(?:PREVIEW-)?(?:NOT-FOR-IMPORT_|IMPORT-ME_)?", "", fn)
def readme_totals(cnt, to_check=None, outside=0, notin=None):
    """v0.4.8 #8: plain totals: 'In the file: 3 new contacts (1 to check). Not in the file: 1 left out, 2 to decide.'
    to_check = rows of the file to check (None: the dry run, where cards to decide are counted instead)."""
    from . import mctrade_export as mx
    t = f"In the file: {plural(cnt[mx.NEW], 'new contact')}" + (f", {plural(cnt[mx.UPDATE], 'update')}" if cnt[mx.UPDATE] else "") \
        + (f" ({to_check} to check)" if to_check else "") + "."
    if notin is None:
        notin = [x for x in (f"{cnt[mx.LEFT]} left out" if cnt[mx.LEFT] else "", f"{cnt[mx.REVIEW]} to decide" if to_check is None and cnt[mx.REVIEW] else "",
                             f"{outside} to decide" if outside else "") if x]
    return t + (" Not in the file: " + ", ".join(notin) + "." if notin else "")

def comp_name(fn, suf):
    """Companion file name. The UPDATES file never says '-new-' (it holds updates of EXISTING profiles, not new ones)."""
    base = os.path.splitext(fn)[0]
    if suf == "-UPDATES.xlsx": base = base.replace("-MCTrade-new-", "-MCTrade-")
    # v0.4.6 #7: the import files say IMPORT-ME, every other file says NOT-FOR-IMPORT (older exports keep their names)
    if base.startswith("PREVIEW-"): return preview_name(base) + suf
    if base.startswith("IMPORT-ME_"): return ("IMPORT-ME-UPDATES_" if suf == "-UPDATES.xlsx" else "NOT-FOR-IMPORT_") + base[len("IMPORT-ME_"):] + suf
    return base + suf

# round 15 (second-round bots): only ONE real export of a batch (or of overlapping cards) runs at a time. A second one
# started while the first is still building gets 409 'export_busy' instead of exporting (and marking) the same cards twice.
EXPORT_BUSY = {}; _EXL = threading.Lock()

def _export_claim(key, ids, who):
    with _EXL:
        for k, (ids2, w2, t2) in EXPORT_BUSY.items():
            if k == key or (ids & ids2): return {"by": w2, "since": t2}
        EXPORT_BUSY[key] = (ids, who, time.time()); return None
def _export_release(key):
    with _EXL: EXPORT_BUSY.pop(key, None)
def export_versions(sel):
    """Round 16 (manager/coordinator r15 H1): what an export sign-off was reviewed against: the template, the contacts file
    and the exact cards (id + version). The page sends these back with the real export; if any changed meanwhile the
    export is refused (409 'stale_review') instead of building a file nobody looked at."""
    # v0.4.12 (review F6): plus the card settings that decide which file a row goes to, or whether it is held back
    ms = U.all_metas(); mv = lambda i: "|".join(str((ms.get(i) or {}).get(k) or "") for k in ("crm_action", "kept_apart", "needs_admin"))
    cv = hashlib.sha256(",".join(f"{c['id']}:{U.card_version(c)}:{mv(c['id'])}" for c in sorted(sel, key=lambda c: c["id"])).encode()).hexdigest()[:12]
    return {"template": U.template_version(), "contacts": U.contacts_version(), "cards": cv}
SUSPECT_WEB = re.compile(r"phish|login|verify|secure-?|account|update-?info|signin|wallet|bit\.ly|tinyurl|xn--|^\d+\.\d+\.\d+\.\d+$", re.I)
def _wdom(w):
    w = re.sub(r"^[a-z]+://", "", str(w or "").strip().lower()).split("/")[0].split(":")[0]
    return w[4:] if w.startswith("www.") else w
def website_warnings(sel, metas, rows=None, merged=None):
    """Round 16 (mctrade r15 M2): a website that does not match the card's email domain(s) AND either came in through a
    merge or looks like a phishing / login / short-link / IP address domain. -> [(card id, website, why)] for the READ-ME."""
    merged_into = {}
    for cid_, m_ in (metas or {}).items():
        d_ = str((m_ or {}).get("dup_resolution") or "")
        if d_.startswith("merged_into:"):
            try: merged_into.setdefault(int(d_.split(":", 1)[1]), []).append(cid_)
            except ValueError: pass
    for tgt, ids_ in (merged or {}).items():   # round 17 (mctrade r16): rows merged at export time count as merges too
        for x_ in ids_:
            if x_ not in merged_into.get(tgt, []): merged_into.setdefault(tgt, []).append(x_)
    out = []
    # round 17 (mctrade r16): with the export's rows, only cards that ARE in an import file are listed, under the exported
    # row's card id, with the website the row really carries (it may have come from a merged-away card)
    items = [(c, str(r_.get("Website") or r_.get("IndWebsite") or "")) for c, r_ in rows] if rows is not None else [(c, c["fields"].get("website") or "") for c in sel]
    for c, w in items:
        w = (w or "").strip(); d = _wdom(w)
        if not d: continue
        eds = {e.split("@")[-1].strip().lower() for e in c["fields"].get("emails") or [] if "@" in e}
        if any(d == e or d.endswith("." + e) or e.endswith("." + d) for e in eds): continue
        why = []
        if SUSPECT_WEB.search(d): why.append("the domain looks like a phishing, login or short-link address")
        if c["id"] in merged_into: why.append("it came in through a merge (from #" + ", #".join(str(x) for x in merged_into[c["id"]]) + ")")
        if why: out.append((c["id"], w, "; ".join(why) + (f"; it doesn't match the email domain {', '.join(sorted(eds))}" if eds else "")))
    return out
def waiting_ids(batch=""):
    """v0.4.7 #4 / v0.4.8 #8: pending cards waiting for an admin in this batch (all batches when none is given)."""
    cs_, m_ = U._active()
    return sorted(c["id"] for c in cs_ if c["status"] == "pending" and (m_.get(c["id"]) or {}).get("needs_admin")
                  and (not batch or ((m_.get(c["id"]) or {}).get("batch") or "__none__") == batch))
@U.router.get("/api/export/summary")
def export_summary(request: Request, status: str = "approved_new", batch: str = ""):
    """v0.4 #5: the counts for the export sheet ("12 new, 1 to check"), from the same server re-check the export runs.
    Counts only: no names, no file, nothing marked as exported."""
    if _filter_refused(status, request.state.user):
        return JSONResponse({"error": "Only an admin can export pending, rejected or all cards.", "admin_only": True}, 403)
    if _batch_refused(batch, request.state.user): return JSONResponse({"error": BATCH_MSG, "admin_only": True}, 403)
    sel, metas, probs = _export_set(status, batch)
    w_ = waiting_ids(batch); u_ = request.state.user or {}
    # v0.4.8 #12: an admin also sees how many wait in ALL batches (one list across batches)
    na_all = len(waiting_ids("")) if u_.get("role") == "admin" else None
    sc_ = counts_for(sel, metas, status, _inc(request), waiting=w_ if batch else None)
    return {"blocked": bool(probs), "needs_admin": len(w_), "needs_admin_all": na_all, "why": why_line(sc_, len(sel))} | sc_

def summary_counts(sel, metas, status, include_changed=False, waiting=None):
    """v0.4.2 #4: the one count source for the Export page line, the sheet, the README and the CSV Check column."""
    out = {"count": len(sel), "new": 0, "updated": 0, "to_check": 0, "left_out": 0, "need_decision": 0, "check_ids": [], "changed_ids": []}
    if sel:
        res = build_export("xlsx", sel, metas, status, "summary", "", preview=True, include_changed=include_changed, waiting=waiting)
        sm = res[4] if len(res) > 4 else {}
        view = sm.get("_view") or {}
        # v0.4.1 #4 (mctrade-admin H1): the view's decisions have STRING keys ('647'); the left-out ids are numbers
        dz = {str(k): v for k, v in (sm.get("decisions") or view.get("decisions") or {}).items()}; left = view.get("left") or []
        rv = {x[0] for x in left if dz.get(str(x[0])) == "review"}
        chk = set(sm.get("check_ids") or [])
        out["ready"] = sm.get("ready")
        rd = sm.get("ready") or {}; inf = set(sm.get("in_file_ids") or [])
        # v0.4.9 #1: 'to check' = rows IN the file to check (the Ready line's number); cards to decide are counted apart
        out |= {"new": len(view.get("rows") or []), "updated": len(view.get("updates") or []), "to_check": rd.get("to_check", 0), "left_out": len(left) - len(rv),
                "need_decision": rd.get("outside", 0), "check_ids": sorted(i for i in set(sm.get("to_check_ids") or []) if i in inf), "decide_ids": sorted(rv), "changed_ids": sm.get("changed_ids") or [],
                "repeat_ids": sm.get("repeat_ids") or [], "held_files": sm.get("held_files") or {}, "rows": (sm.get("ready") or {}).get("rows", 0),
                "changed_file": next((x for x in ((sm.get("held_files") or {}).get(str(i)) for i in (sm.get("changed_ids") or [])[:5]) if x), "")}
    return out

_RC = {}; _RCL = threading.Lock()
def counts_for(sel, metas, status, include_changed=False, waiting=None):
    """v0.4.9 #1: THE count source (batch list, Export page, Save dialog, README, SUMMARY, dry run): summary_counts,
    cached per exact card set + template + contacts + earlier import files + cards waiting for an admin."""
    v = export_versions(sel); sig = _IDX_CACHE["sig"] or tuple(U.db().execute("SELECT COUNT(*), MAX(id) FROM exports WHERE kind='xlsx'").fetchone())
    # card metadata (kept apart, needs admin, re-export ...) and the newest history row (cross-batch confirms) are part of the key
    mh = hashlib.sha256(json.dumps([metas.get(c["id"]) for c in sorted(sel, key=lambda c: c["id"])], default=str, sort_keys=True).encode()).hexdigest()[:12]
    hmax = U.db().execute("SELECT MAX(id) FROM card_history").fetchone()[0]
    k = (status, v["cards"], v["template"], v["contacts"], tuple(sig), tuple(waiting or []), bool(include_changed), len(sel), mh, hmax)
    with _RCL:
        if k in _RC: return _RC[k]
    out = summary_counts(sel, metas, status, include_changed, waiting=waiting)
    with _RCL:
        if len(_RC) > 300: _RC.clear()
        _RC[k] = out
    return out
def why_line(sc, count):
    """v0.4.9 #1 / data-manager: one short line saying why nothing (or less) goes into the import file, and what to do."""
    r = (sc or {}).get("ready") or {}
    if not count: return "Nothing to export yet. Approve some cards first."
    if r.get("rows"): 
        p = notin_parts(r)
        return ("Not in the file: " + ", ".join(p) + ". See the Left-out list in the Save dialog." ) if p else ""
    held = (r.get("changed") or 0) + (r.get("repeat") or 0)
    if held and held == count:
        return f"{plural(held, 'card')} already exported. Nothing new to import."
    p = notin_parts(r)
    return ("Nothing new to import: " + ", ".join(p) + ". Open those cards to fix or decide them.") if p else "Nothing new to import."
def _inc(request):
    """v0.4.7 #2: an admin may put cards exported before back in the import file as new (the dialog's option)."""
    return request.query_params.get("include_changed") == "1" and (request.state.user or {}).get("role") == "admin"

MT = {"csv": "text/csv", "vcf": "text/vcard", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
@U.router.get("/api/export/{kind}")
def do_export(kind: str, request: Request, status: str = "approved_new", batch: str = "", preview: bool = False):
    """preview=1: build and return the file(s) but change nothing (no 'exported' marks, no history, no stored copy)."""
    if kind == "dryrun": return U.export_dryrun(request, status, batch)   # same URL family; defined at the end of this module
    if kind not in MT: raise HTTPException(404)
    from . import demo
    if demo.enabled(): preview = True   # public demo: every export (incl. the MC Trade Excel) is a dry-run preview
    if _filter_refused(status, request.state.user):
        return JSONResponse({"error": "Only an admin can export pending, rejected or all cards. Reviewers export approved cards."}, 403)
    if _batch_refused(batch, request.state.user): return JSONResponse({"error": BATCH_MSG, "admin_only": True}, 403)
    if preview and not batch and request.state.user.get("role") != "admin":   # skeptic R6-1: no all-batches contact dump for reviewers
        return JSONResponse({"error": "Choose one batch under “Batch” to preview. Only admins can preview all batches."}, 403)
    sel, metas, probs = _export_set(status, batch)
    if probs and not preview:   # a preview changes nothing, so look-alikes don't block it
        msg = ("Some cards in this batch look like the same person. Resolve them in the review screen before exporting." if batch else
               "With “All batches”, look-alikes from anyone's batch block the export. Choose your batch under “Batch”, "
               "or ask an admin to resolve the cards listed.")
        return JSONResponse({"error": msg, "duplicates": probs}, 409)
    if not sel: return JSONResponse({"error": "No cards to export."}, 404)
    user = request.state.user
    if not preview:
        seen_ = {k: request.query_params.get("seen_" + k) for k in ("template", "contacts", "cards")}
        _cl = ": cards " + ", ".join(f"#{c['id']}" for c in sel[:500])   # round 17 (manager r16 L-1): refused rows name the cards
        if not all(seen_.values()):   # round 17 (manager r16 M-1, skeptic L1): the sign-off versions are REQUIRED for a real export
            miss = [k for k, v in seen_.items() if not v]
            auth.alog(user["username"], f"refused export ({kind}, " + (f"batch {batch}" if batch else "all batches") + "): sent without the sign-off versions (" + ", ".join("seen_" + k for k in miss) + "); nothing was exported" + _cl)
            return JSONResponse({"error": "Nothing was exported: the sign-off check is missing. Export from the Export page.",
                                 "need": "version_required", "missing": ["seen_" + k for k in miss]}, 428)   # round 21 (skeptic r20 L1): no versions in a 428
        if True:
            now_ = export_versions(sel); drift = [k for k, v in seen_.items() if v and v != now_[k]]
            if drift:
                words = {"template": "the MC Trade template", "contacts": "the MC Trade contacts file", "cards": "the cards (edited, approved or rejected)"}
                auth.alog(user["username"], f"refused export ({kind}, " + (f"batch {batch}" if batch else "all batches") + "): changed since the sign-off was shown: " + ", ".join(words[k] for k in drift) + "; nothing was exported" + _cl)
                return JSONResponse({"error": "Since you opened the export sign-off, " + " and ".join(words[k] for k in drift) + " changed. Nothing was exported: check the sign-off again, then export.",
                                     "need": "stale_review", "changed": drift, "now": now_}, 409)
        xkey = "batch:" + (batch or "__all__"); busy_ = _export_claim(xkey, {c["id"] for c in sel}, user["username"])
        if busy_: auth.alog(user["username"], f"refused export ({kind}, " + (f"batch {batch}" if batch else "all batches") + f"): another export of these cards was running (started by {busy_['by']} at {time.strftime('%H:%M:%S', time.localtime(busy_['since']))}); nothing was exported" + _cl)
        if busy_: return JSONResponse({"error": f"Another export of these cards is running right now (started by {busy_['by']} at {time.strftime('%H:%M:%S', time.localtime(busy_['since']))}). Nothing was exported: wait for it to finish, then check the export history before exporting again.", "need": "export_busy"}, 409)
        try: return _do_export_real(kind, request, status, batch, sel, metas, user, seen_)
        finally: _export_release(xkey)
    blabel = next(((metas.get(c["id"]) or {}).get("batch_label") or "" for c in sel), "") if batch else ""
    res = build_export(kind, sel, metas, status, f"{user['display']} ({user['initials']})", blabel, preview=preview, include_changed=_inc(request), waiting=waiting_ids(batch) if batch else None, vis=U.visible_ids(user), viewer=user)
    body, mt, fn = res[:3]; extras = res[3] if len(res) > 3 else {}
    view = res[4].pop("_view", None) if len(res) > 4 else None
    if len(res) > 4: res[4].pop("people", None)
    summ_h = {"X-Summary": json.dumps(res[4])} if len(res) > 4 else {}
    if preview and view is not None and request.query_params.get("view") == "1":   # on-screen preview (no file)
        auth.alog(user["username"], f"viewed the PREVIEW of the MC Trade files on screen ({len(sel)} {U._s(len(sel), 'card')}" + (f", batch {blabel or batch}" if batch else ", all batches") + "): " + ", ".join(f"#{c['id']}" for c in sel[:500]))
        return view | res[4] | {"versions": export_versions(sel)}
    if preview:
        fn = preview_name(fn)
        if request.query_params.get("companion") and request.query_params["companion"] not in extras:   # mctrade-r5 F1
            return JSONResponse({"error": "There is no such file in this export (for example no update rows, so no UPDATES file)."}, 404)
        nm_ = request.query_params.get("companion") or "main file"
        auth.alog(user["username"], f"downloaded the PREVIEW {kind} ({nm_}) of {len(sel)} {U._s(len(sel), 'card')}, {FILTERS.get(status, status)}"
                  + (f", batch {blabel or batch}" if batch else ", all batches") + ": " + ", ".join(f"#{c['id']}" for c in sel[:500]) + ("…" if len(sel) > 500 else ""))
        hdrs = {"Content-Disposition": f'attachment; filename="{fn}"', "X-Card-Count": str(len(sel)), "X-Preview": "1", "X-Companions": ",".join(extras)} | summ_h
        if request.query_params.get("companion"):
            nm = request.query_params["companion"]; mt2, suf = COMPANION[nm]
            return Response(extras[nm], media_type=mt2, headers=hdrs | {"Content-Disposition": f'attachment; filename="{comp_name(fn, suf)}"'})
        return Response(body, media_type=mt, headers=hdrs)
    raise AssertionError('unreachable')

def _stale_export(kind, batch, user, sel, changed, now_, when):
    words = {"template": "the MC Trade template", "contacts": "the MC Trade contacts file", "cards": "the cards (edited, approved or rejected)"}
    auth.alog(user["username"], f"refused export ({kind}, " + (f"batch {batch}" if batch else "all batches") + f"): changed {when}: "
              + ", ".join(words[k] for k in changed) + "; nothing was exported: cards " + ", ".join(f"#{c['id']}" for c in sel[:500]))
    return JSONResponse({"error": f"{' and '.join(words[k] for k in changed).capitalize()} changed {when}. Nothing was exported: check the sign-off again, then export.",
                         "need": "stale_review", "changed": changed, "now": now_}, 409)

def _do_export_real(kind, request, status, batch, sel, metas, user, seen=None):
    # re-read the set INSIDE the claim: cards exported by a run that just finished are not exported again by accident
    sel, metas, probs = _export_set(status, batch)
    if probs: return JSONResponse({"error": "Some cards in this batch look like the same person. Resolve them in the review screen before exporting.", "duplicates": probs}, 409)
    if not sel: return JSONResponse({"error": "No cards to export (they may have just been exported by someone else: see the export history)."}, 404)
    # audit 2026-10 P0-2 / F01: the re-read set must still be exactly what the sign-off showed (a card approved in between
    # would otherwise be exported without anyone having looked at it on the sign-off)
    built_v = export_versions(sel)
    if seen:
        drift = [k for k, v in seen.items() if v and v != built_v[k]]
        if drift: return _stale_export(kind, batch, user, sel, drift, built_v, "since the sign-off was shown")
    blabel = next(((metas.get(c["id"]) or {}).get("batch_label") or "" for c in sel), "") if batch else ""
    res = build_export(kind, sel, metas, status, f"{user['display']} ({user['initials']})", blabel, include_changed=_inc(request), waiting=waiting_ids(batch) if batch else None, vis=U.visible_ids(user), viewer=user)
    body, mt, fn = res[:3]; extras = res[3] if len(res) > 3 else {}
    if len(res) > 4: res[4].pop("_view", None)
    people_ = res[4].pop("people", []) if len(res) > 4 else []   # stored with the export, not sent back in a header
    infile = set(res[4]["in_file_ids"]) if len(res) > 4 else {c["id"] for c in sel}
    if kind == "xlsx" and len(res) > 4 and not res[4]["ready"]["rows"]:
        # v0.4.8 #1: nothing new to import: no IMPORT-ME file, no export row, nothing marked exported
        auth.alog(user["username"], f"export ({kind}, " + (f"batch {blabel or batch}" if batch else "all batches") + "): nothing new to import, no file made: " + res[4]["ready"]["text"])
        return JSONResponse({"nothing": True, "message": "Nothing new to import.", "ready": res[4]["ready"],
                             "changed_ids": res[4]["changed_ids"], "repeat_ids": res[4]["repeat_ids"]})
    msel = [c for c in sel if c["id"] in infile]
    summ_h = {"X-Summary": json.dumps(res[4])} if len(res) > 4 else {}
    stored = uuid.uuid4().hex + "_" + fn
    with open(os.path.join(U.EXP, stored), "wb") as fh: fh.write(body)
    for nm, data in extras.items():
        with open(os.path.join(U.EXP, stored + "." + nm), "wb") as fh: fh.write(data)
    now = time.time()
    # audit 2026-10 P0-2: mark the cards only if nothing changed while the file was being built (an edit made meanwhile
    # is not in the file; marking it exported would hide it from the next export). Checked and written under WRITE_LOCK,
    # in one transaction, so no save can slip in between; the export id comes from this insert (F02), not MAX(id).
    with U.WRITE_LOCK:
        cur = [store.get(c["id"]) for c in sel]
        now_v = export_versions([x for x in cur if x]) if all(cur) else None
        if now_v != built_v:
            for p_ in [stored] + [stored + "." + nm for nm in extras]:
                try: os.remove(os.path.join(U.EXP, p_))
                except OSError: pass
            return _stale_export(kind, batch, user, sel, [k for k in built_v if not now_v or now_v[k] != built_v[k]], now_v, "while the file was being built")
        curm = U.all_metas(); tv = U.template_version()
        with U.db() as c:
            # v0.4.12 (review F1): card_ids = the cards IN the import file; listed_ids = every card the export's files mention
            # (the extra-columns CSV, READ-ME and SUMMARY also list the left-out ones), so delete-on-request finds every copy
            eid = c.execute("INSERT INTO exports(ts,username,initials,kind,filter,count,filename,card_ids,stored,batch_label,people,listed_ids) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                            (now, user["username"], user["initials"], kind, status, len(msel), fn, json.dumps([x["id"] for x in msel]), stored, blabel if batch else "",
                             json.dumps(people_) if kind == "xlsx" and len(res) > 4 else None,
                             json.dumps(sorted({x["id"] for x in sel} | {x["id"] for x in msel})))).lastrowid
            for x in msel:   # v0.4.8 #1: only the cards really in the file are marked exported
                m = curm.get(x["id"]) or {}
                c.execute("INSERT INTO card_meta(card_id) VALUES(?) ON CONFLICT(card_id) DO NOTHING", (x["id"],))
                if not m.get("exported"): c.execute("UPDATE card_meta SET exported=? WHERE card_id=?", (now, x["id"]))
                # v0.4.9 #2: only an MC Trade import file is 'sent to MC Trade'. A contacts list (CSV/vCard) never clears
                # 'Edited after export' and is logged as its own kind of export.
                if kind == "xlsx" and m.get("reexport"): c.execute("UPDATE card_meta SET reexport=NULL WHERE card_id=?", (x["id"],))
                c.execute("INSERT INTO card_history(card_id,ts,reviewer,action,from_status,to_status,detail) VALUES(?,?,?,?,?,?,?)",
                          (x["id"], now, user["initials"], "exported" if kind == "xlsx" else "exported_list", None, None,
                           json.dumps({"file": fn, "template_version": tv} if kind == "xlsx" else {"file": fn, "kind": kind})))
    return Response(body, media_type=mt, headers={"Content-Disposition": f'attachment; filename="{fn}"', "X-Card-Count": str(len(msel)),
                                                  "X-Export-Id": str(eid), "X-Companions": ",".join(extras)} | summ_h)

COMPANION = {"extra.csv": ("text/csv", "-extra-columns.csv"), "readme.txt": ("text/plain", "-READ-ME.txt"),
             "updates.xlsx": (MT["xlsx"], "-UPDATES.xlsx"), "summary.xlsx": (MT["xlsx"], "-SUMMARY.xlsx")}
@U.router.get("/api/exports/{eid}/companion/{name}")
def export_companion(eid: int, name: str, request: Request):
    """Companion files of an MC Trade import file: the non-template columns + rows left out, and the READ ME notes."""
    if name not in COMPANION: raise HTTPException(404)
    r = U.db().execute("SELECT * FROM exports WHERE id=?", (eid,)).fetchone() if 0 < eid < 2**63 else None   # fuzz 2026-10: a huge id was a 500
    u = request.state.user
    if r and u["role"] != "admin" and r["username"] != u["username"]:
        return JSONResponse({"error": "Only the person who made this export (or an admin) can download it."}, 403)
    p = os.path.join(U.EXP, (r["stored"] or "") + "." + name) if r and r["stored"] else ""
    if not p or not os.path.isfile(p): return JSONResponse({"error": "No such file (it may have been removed by the retention policy)."}, 404)
    mt, suf = COMPANION[name]
    ids_ = json.loads(r["card_ids"] or "[]")
    auth.alog(u["username"], f"downloaded the {suf.strip('-')} companion of export #{eid} ({r['filename']}; cards " + ", ".join(f"#{i}" for i in ids_[:500]) + ("…" if len(ids_) > 500 else "") + ")")
    return FileResponse(p, media_type=mt, filename=comp_name(r["filename"], suf))

@U.router.get("/api/exports")
def export_history(request: Request):
    out = []; u = request.state.user
    for r in U.db().execute("SELECT * FROM exports ORDER BY id DESC LIMIT 200"):
        d = dict(r); d["available"] = bool(d["stored"]) and os.path.isfile(os.path.join(U.EXP, d["stored"])); d["filter_text"] = FILTERS.get(d["filter"], d["filter"])
        d["own"] = d["username"] == u["username"]; d["mine"] = u["role"] == "admin" or d["own"]; d["available"] = d["available"] and d["mine"]
        if not d["mine"]: continue   # v0.4.3 #8 (mctrade-admin): a reviewer sees only their own export history
        d["companions"] = [n for n in COMPANION if d["available"] and os.path.isfile(os.path.join(U.EXP, d["stored"] + "." + n))]
        d["card_ids"] = json.loads(d["card_ids"] or "[]"); d.pop("stored"); out.append(d)
    return out

@U.router.get("/api/exports/{eid}/download")
def export_again(eid: int, request: Request):
    r = U.db().execute("SELECT * FROM exports WHERE id=?", (eid,)).fetchone() if 0 < eid < 2**63 else None   # fuzz 2026-10: a huge id was a 500
    if not r or not r["stored"] or not os.path.isfile(os.path.join(U.EXP, r["stored"])):
        return JSONResponse({"error": "That file is no longer stored. Export again."}, 404)
    u = request.state.user
    if u["role"] != "admin" and r["username"] != u["username"]:
        return JSONResponse({"error": "Only the person who made this export (or an admin) can download it again."}, 403)
    ids = json.loads(r["card_ids"] or "[]")
    auth.alog(request.state.user["username"], f"re-downloaded export #{eid} ({r['filename']}; {len(ids)} {U._s(len(ids), 'card')}: " + ", ".join(f"#{i}" for i in ids[:500]) + (" …" if len(ids) > 500 else "") + ")")
    for i in ids: U.log(i, request.state.user["initials"], "redownloaded", detail={"export_id": eid, "file": r["filename"]})
    return FileResponse(os.path.join(U.EXP, r["stored"]), media_type=MT.get(r["kind"]), filename=r["filename"])
